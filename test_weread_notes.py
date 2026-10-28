"""离线回归测试：不读环境变量中的凭证，不访问网络，不写入真实 Notion。"""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from urllib.request import HTTPSHandler
from zoneinfo import ZoneInfo

import weread_notes as w


class ScriptedHTTPS(HTTPSHandler):
    """走真实 urllib/CookieJar 管线，但将 HTTPS 传输替换为内存响应。"""
    def __init__(self, responses):
        super().__init__()
        self.responses = list(responses)
        self.requests = []

    def https_open(self, request):
        from email.message import Message
        from urllib.response import addinfourl
        self.requests.append((request.full_url, request.get_method(), dict(request.header_items())))
        if not self.responses:
            raise AssertionError("Unexpected extra request")
        status, body, cookie = self.responses.pop(0)
        headers = Message()
        headers.add_header("Content-Type", "application/json")
        if cookie:
            headers.add_header("Set-Cookie", cookie)
        response = addinfourl(io.BytesIO(w.json.dumps(body).encode()), headers, request.full_url, status)
        response.msg = "Test response"
        return response


class SessionTests(unittest.TestCase):
    def client(self, responses):
        client = w.Weread("wr_vid=123; wr_skey=old-example; wr_rt=other-example")
        transport = ScriptedHTTPS(responses)
        client.opener = w.build_opener(w.HTTPCookieProcessor(client.cookiejar), w.NoRedirect(), transport)
        return client, transport

    @patch("weread_notes.time.sleep")
    def test_rotated_cookie_used_on_next_request_without_duplicates(self, _):
        client, transport = self.client([
            (200, {"ok": 1}, "wr_skey=new-example; Domain=.weread.qq.com; Path=/; Secure"),
            (200, {"ok": 1}, None),
        ])
        client.request("/web/first")
        client.request("/web/second")
        sent = transport.requests[1][2]["Cookie"]
        self.assertIn("wr_skey=new-example", sent)
        self.assertNotIn("old-example", sent)
        self.assertEqual(sent.count("wr_skey="), 1)
        self.assertIn("wr_rt=other-example", sent)

    @patch("weread_notes.time.sleep")
    def test_session_refresh_then_retry_once(self, _):
        client, transport = self.client([
            (200, {"errcode": -2012}, None),
            (200, {}, "wr_skey=refreshed-example; Path=/; Secure"),
            (200, {"books": []}, None),
        ])
        self.assertEqual(client.shelf(), [])
        self.assertEqual([r[1] for r in transport.requests], ["GET", "HEAD", "GET"])
        self.assertIn("wr_skey=refreshed-example", transport.requests[2][2]["Cookie"])

    @patch("weread_notes.time.sleep")
    def test_rejected_shelf_with_valid_notebook_distinguishes_endpoint(self, _):
        client, transport = self.client([
            (200, {"errcode": -2012}, None), (200, {}, None),
            (200, {"errcode": -2012}, None), (200, {"books": []}, None),
        ])
        with self.assertRaisesRegex(w.ExportError, "个人笔记接口验证成功"):
            client.shelf()
        self.assertEqual(len(transport.requests), 4)

    @patch("weread_notes.time.sleep")
    def test_failed_probe_does_not_recursively_refresh(self, _):
        client, transport = self.client([
            (200, {"errcode": -2012}, None), (200, {}, None),
            (200, {"errcode": -2012}, None), (200, {"errcode": -2012}, None),
        ])
        with self.assertRaisesRegex(w.ExportError, "个人登录状态未能验证"):
            client.shelf()
        self.assertEqual(len(transport.requests), 4)

    @patch("weread_notes.time.sleep")
    def test_http401_receives_cookie_and_recovers(self, _):
        client, transport = self.client([
            (401, {}, "wr_skey=new-example; Domain=.weread.qq.com; Path=/; Secure"),
            (200, {}, None), (200, {"books": []}, None),
        ])
        self.assertEqual(client.shelf(), [])
        self.assertIn("wr_skey=new-example", transport.requests[1][2]["Cookie"])

    @patch("weread_notes.time.sleep")
    def test_agent_never_receives_web_cookies(self, _):
        client, web = self.client([])
        client.api_key = "example-not-real-key"
        agent = ScriptedHTTPS([(200, {"books": []}, None)])
        client.agent_opener = w.build_opener(w.NoRedirect(), agent)
        self.assertEqual(client.shelf(), [])
        self.assertFalse(web.requests)
        self.assertNotIn("Cookie", agent.requests[0][2])
        self.assertEqual(agent.requests[0][2]["Authorization"], "Bearer example-not-real-key")

    @patch("weread_notes.time.sleep")
    def test_error_in_data_wrapper_is_not_mistaken_for_missing_notes(self, _):
        client, transport = self.client([(200, {"data": {"errcode": -999}}, None)])
        with self.assertRaisesRegex(w.ExportError, "代码 -999"):
            client.request("/web/book/bookmarklist")

    @patch("weread_notes.time.sleep")
    def test_empty_web_response_falls_back_to_official_legacy_host(self, _):
        mark = {"bookmarkId": "h", "markText": "test quote", "type": 1}
        client, transport = self.client([
            (200, {}, None),
            (200, {"books": [{"bookId": "b", "noteCount": 1}]}, None),
            (200, {}, None),
            (200, {"updated": [mark]}, None),
        ])
        self.assertEqual(client.highlights("b")[0], [mark])
        self.assertIn("type=1", transport.requests[2][0])
        self.assertTrue(transport.requests[3][0].startswith("https://i.weread.qq.com/book/bookmarklist?"))
        self.assertIn("wr_skey=old-example", transport.requests[3][2]["Cookie"])
        self.assertNotIn("Authorization", transport.requests[3][2])

    @patch("weread_notes.time.sleep")
    def test_empty_web_type_variant_can_succeed_without_legacy(self, _):
        mark = {"bookmarkId": "h", "markText": "test quote", "type": 1}
        client, transport = self.client([
            (200, {}, None),
            (200, {"books": [{"bookId": "b", "noteCount": 1}]}, None),
            (200, {"updated": [mark]}, None),
        ])
        self.assertEqual(client.highlights("b")[0], [mark])
        self.assertEqual(len(transport.requests), 3)

    @patch("weread_notes.time.sleep")
    def test_empty_fallback_not_success_when_overview_says_nonempty(self, _):
        client, transport = self.client([
            (200, {}, None),
            (200, {"books": [{"bookId": "b", "noteCount": 1}]}, None),
            (200, {"updated": []}, None),
            (200, {"updated": []}, None),
        ])
        with self.assertRaisesRegex(w.ExportError, "已停止导出"):
            client.highlights("b")

    def test_legacy_paths_are_restricted(self):
        client = w.Weread("wr_vid=123; wr_skey=example")
        with self.assertRaises(w.ExportError):
            client.request("/arbitrary/path", legacy=True)
        with self.assertRaises(w.ExportError):
            client.request("/book/bookmarklist", legacy=True, agent=True)


class PersonalExportTests(unittest.TestCase):
    def client(self, **kwargs):
        client = w.Weread(api_key="wrk-example-only", **kwargs)
        def request(path, params=None, **options):
            if path == "/book/info":
                return {"title": "测试书"}
            if path == "/book/chapterinfo":
                return {"chapters": [{"chapterUid": 1, "title": "第一章"}]}
            if path == "/book/bookmarklist":
                return {"updated": [{"bookmarkId": "h", "markText": "个人划线", "chapterUid": 1}]}
            raise AssertionError("Unexpected request")
        client.request = request
        def pages(path, *args, **options):
            if path == "/review/list/mine":
                return [{"review": {"reviewId": "r", "content": "个人想法", "chapterUid": 1}}]
            raise w.ExportError("测试：点赞接口不可访问")
        client.pages = Mock(side_effect=pages)
        client.liked_thoughts = Mock(side_effect=w.ExportError("测试：点赞接口不可访问"))
        return client

    def test_personal_only_uses_no_cookie_or_public_review_endpoint(self):
        client = self.client(personal_only=True)
        book = client.collect({"bookId": "b"}, ZoneInfo("UTC"), 0)
        self.assertEqual(client.cookie, "")
        client.pages.assert_called_once()
        self.assertEqual(client.pages.call_args.args[0], "/review/list/mine")
        self.assertTrue(any("未抓取已点赞他人评论" in notice for notice in book.warnings))
        texts = "\n".join(text for kind, text in w.entries(book))
        self.assertIn("个人划线", texts)
        self.assertIn("个人想法", texts)

    @patch("weread_notes.choice", return_value="1")
    def test_partial_export_requires_explicit_choice_and_marks_missing_comments(self, choice):
        client = self.client(interactive_auth=True)
        book = client.collect({"bookId": "b"}, ZoneInfo("UTC"), 0)
        choice.assert_called_once()
        self.assertTrue(any("接口读取失败" in notice for notice in book.warnings))
        self.assertFalse(any("评论接口返回空列表" in notice for notice in book.warnings))

    @patch("weread_notes.choice", return_value="2")
    def test_declining_partial_export_stops(self, choice):
        with self.assertRaisesRegex(w.ExportError, "用户未选择"):
            self.client(interactive_auth=True).collect({"bookId": "b"}, ZoneInfo("UTC"), 0)

    def test_noninteractive_default_does_not_silently_skip_failed_comments(self):
        with self.assertRaisesRegex(w.ExportError, "点赞接口不可访问"):
            self.client().collect({"bookId": "b"}, ZoneInfo("UTC"), 0)

    @patch("weread_notes.export_word", return_value=Path("demo-not-created.docx"))
    @patch("weread_notes.select_book", return_value={"bookId": "b", "title": "测试书"})
    @patch("weread_notes.secret", return_value="wrk-example-only")
    @patch("weread_notes.choice", side_effect=["2", "1", "2"])
    def test_personal_only_cli_does_not_prompt_for_cookie(self, choice, secret, select, export):
        mock_client = SimpleNamespace(shelf=Mock(return_value=[]), collect=Mock(return_value=w.demo_book(ZoneInfo("UTC"))))
        with patch.object(w.sys, "argv", ["weread_notes.py", "--personal-only"]), patch(
                "weread_notes.Weread", return_value=mock_client) as constructor:
            self.assertEqual(w.main(), 0)
        secret.assert_called_once()
        self.assertEqual(secret.call_args.args[1], "WEREAD_API_KEY")
        self.assertEqual(constructor.call_args.args[0], "")
        self.assertTrue(constructor.call_args.kwargs["personal_only"])


class LikedThoughtTests(unittest.TestCase):
    def test_official_detail_does_not_let_wrapper_state_override_actual_review(self):
        data = {"reviewId": "456_a", "isLiked": 0, "userVid": "123", "review": {
            "content": "actual", "isLike": 1, "userVid": "456", "author": {"userVid": "456"}}}
        row = w.unwrap_review_detail(data)
        self.assertIs(w.like_state(row, official_detail=True), True)
        self.assertNotIn("isLiked", row)
        self.assertEqual(w.author_id(row), "456")
        row = w.unwrap_review_detail({"review": {"content": "x", "isLike": 1, "isLiked": 0}})
        self.assertIsNone(w.like_state(row, official_detail=True))

    def test_wrapper_default_no_longer_hides_liked_link(self):
        client = w.Weread(api_key="wrk-example", liked_source="links", review_ids=["456_a"])
        client.request = Mock(return_value={"reviewId": "456_a", "isLiked": 0, "review": {
            "bookId": "b", "content": "actual", "isLike": 1, "author": {"userVid": 456}}})
        rows, _ = client.liked_thoughts("b", [], [], set())
        self.assertEqual([row["reviewId"] for row in rows], ["456_a"])

    def test_verification_budget_prioritizes_later_liked_hint_without_trusting_it(self):
        client = w.Weread(api_key="wrk-example", max_review_checks=1)
        fetched = []
        def request(path, params=None, **kwargs):
            if path == "/book/bestbookmarks":
                return {"items": [{"chapterUid": 1, "range": "1-10"}]}
            if path == "/book/readreviews":
                rows = [{"reviewId": f"456_r{i}", "review": {"isLike": int(i == 4)}} for i in range(5)]
                return {"reviews": [{"range": "1-10", "hasMore": 0, "pageReviews": rows}]}
            fetched.append(params["reviewId"])
            # Even a positive list hint is discarded when fresh detail says unliked.
            return {"reviewId": params["reviewId"], "review": {"content": "x", "bookId": "b", "isLike": 0}}
        client.request = request
        rows, warnings = client.liked_thoughts("b", [{"chapterUid": 1}], [], set())
        self.assertEqual(fetched, ["456_r4"])
        self.assertEqual(rows, [])
        self.assertTrue(any("仅核验前 1 条" in warning for warning in warnings))

    def test_null_detail_fields_do_not_erase_known_location_or_author(self):
        context = {"chapterUid": 1, "range": "1-10", "author": {"userVid": 456}, "abstract": "摘录"}
        merged = w.merge_review_context(context, {"chapterUid": None, "range": "", "author": {}, "content": "正文"})
        self.assertEqual(merged["chapterUid"], 1)
        self.assertEqual(merged["range"], "1-10")
        self.assertEqual(w.author_id(merged), "456")

    def test_unliked_link_reports_exact_reason_without_printing_content(self):
        client = w.Weread(api_key="wrk-example", liked_source="links", review_ids=["456_a"])
        client.request = Mock(return_value={"reviewId": "456_a", "review": {
            "bookId": "b", "content": "private-content", "isLike": 0}})
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            rows, _ = client.liked_thoughts("b", [], [], set())
        self.assertEqual(rows, [])
        self.assertIn("官方详情当前返回未点赞（状态 0）", output.getvalue())
        self.assertNotIn("private-content", output.getvalue())

    def test_like_state_accepts_only_verified_detail_alias(self):
        self.assertIsNone(w.like_state({"isLike": 1}))
        self.assertIs(w.like_state({"isLike": 1}, official_detail=True), True)
        self.assertIs(w.like_state({"isLike": 0}, official_detail=True), False)
        self.assertIsNone(w.like_state({"isLike": 1, "isLiked": 0}, official_detail=True))
        self.assertIsNone(w.like_state({"likesCount": 10000}, official_detail=True))
        self.assertIsNone(w.like_state({"isLike": "unknown"}, official_detail=True))

    def test_links_only_checks_current_like_state_and_deduplicates(self):
        client = w.Weread(api_key="wrk-example", liked_source="links", review_ids=["456_a", "456_b", "456_a"])
        client.request = Mock(side_effect=[
            {"reviewId": "456_a", "review": {"bookId": "b", "content": "liked", "isLike": 1, "author": {"userVid": 456}}},
            {"reviewId": "456_b", "review": {"bookId": "b", "content": "unliked", "isLike": 0, "author": {"userVid": 456}}},
        ])
        rows, warnings = client.liked_thoughts("b", [], [], set())
        self.assertEqual([row["reviewId"] for row in rows], ["456_a"])
        self.assertEqual(rows[0]["isLiked"], 1)
        self.assertEqual(client.request.call_count, 2)
        self.assertTrue(any("仅核验本次提供" in warning for warning in warnings))

    def test_provided_link_from_different_book_is_not_misfiled(self):
        client = w.Weread(api_key="wrk-example", liked_source="links", review_ids=["456_a"])
        client.request = Mock(return_value={"reviewId": "456_a", "review": {"bookId": "other", "isLike": 1, "content": "x"}})
        with self.assertRaisesRegex(w.ExportError, "另一本书"):
            client.liked_thoughts("b", [], [], set())

    def test_own_review_link_is_not_fetched_or_exported(self):
        client = w.Weread(api_key="wrk-example", liked_source="links", review_ids=["123_a"])
        client.request = Mock()
        rows, _ = client.liked_thoughts("b", [], [], {"123_a"})
        self.assertEqual(rows, [])
        client.request.assert_not_called()

    def test_paragraphs_page_then_verify_each_detail(self):
        client = w.Weread(api_key="wrk-example")
        calls = []
        def request(path, params=None, **kwargs):
            calls.append((path, params, kwargs))
            if path == "/book/bestbookmarks":
                self.assertIs(type(params["chapterUid"]), int)
                return {"items": [{"chapterUid": 1, "range": "10-20", "markText": "excerpt"}]}
            if path == "/book/readreviews":
                cursor = params["reviews"][0]["maxIdx"]
                return {"reviews": [{"range": "10-20", "pageReviews": [
                    {"reviewId": "456_a" if cursor == 0 else "456_b", "review": {"content": "sample"}}],
                    "hasMore": 1 if cursor == 0 else 0, "maxIdx": 1, "synckey": 2}]}
            rid = params["reviewId"]
            return {"reviewId": rid, "review": {"bookId": "b", "chapterUid": 1, "content": "x",
                "isLike": 1 if rid == "456_a" else 0, "author": {"userVid": 456}}}
        client.request = request
        rows, warnings = client.liked_thoughts("b", [{"chapterUid": 1}], [], set())
        self.assertEqual([r["reviewId"] for r in rows], ["456_a"])
        self.assertEqual(rows[0]["range"], "10-20")
        self.assertTrue(any("不代表全书完整" in warning for warning in warnings))
        self.assertEqual(len([call for call in calls if call[0] == "/book/readreviews"]), 2)

    def test_scan_stagnation_does_not_report_empty_success(self):
        client = w.Weread(api_key="wrk-example")
        def request(path, params=None, **kwargs):
            if path == "/book/bestbookmarks":
                return {"items": [{"chapterUid": 1, "range": "10-20"}]}
            return {"reviews": [{"range": "10-20", "pageReviews": [], "hasMore": 1, "maxIdx": 0, "synckey": 0}]}
        client.request = request
        with self.assertRaisesRegex(w.ExportError, "不能将此结果视为没有点赞"):
            client.liked_thoughts("b", [{"chapterUid": 1}], [], set())

    def test_known_link_exports_other_reader_without_cookie(self):
        client = w.Weread(api_key="wrk-example", liked_source="links", review_ids=["456_a"])
        def request(path, params=None, **kwargs):
            if path == "/book/info":
                return {"title": "书"}
            if path == "/book/chapterinfo":
                return {"chapters": [{"chapterUid": 1, "title": "第一章"}]}
            if path == "/book/bookmarklist":
                return {"updated": [{"bookmarkId": "h", "chapterUid": 1, "range": "1-2", "markText": "摘录"}]}
            return {"reviewId": "456_a", "review": {"bookId": "b", "chapterUid": 1, "range": "1-2",
                "content": "读者的内容", "isLike": 1, "author": {"userVid": 456, "name": "读者甲"}}}
        client.request = request
        client.pages = Mock(return_value=[{"review": {"reviewId": "123_own", "content": "自己的内容", "author": {"userVid": 123}}}])
        notebook = client.collect({"bookId": "b"}, ZoneInfo("UTC"), 0)
        text = "\n".join(text for _, text in w.entries(notebook))
        self.assertIn("读者观点\n读者的内容", text)
        self.assertNotIn("读者甲", text)
        self.assertIn("读者的内容", text)
        self.assertIn("自己的内容", text)
        self.assertEqual(client.cookie, "")


class ReviewProbeTests(unittest.TestCase):
    def test_reference_requires_review_id_not_book_hash(self):
        self.assertEqual(w.review_id_from_reference("123_Ab-c"), "123_Ab-c")
        self.assertEqual(w.review_id_from_reference("https://weread.qq.com/reviewDetail?reviewId=123_Ab-c"), "123_Ab-c")
        for invalid in ("330075595", "https://weread.qq.com/web/reader/bookhash", "https://example.com/?reviewId=123_Ab-c"):
            with self.assertRaises(w.ExportError):
                w.review_id_from_reference(invalid)

    def test_probe_reports_flags_without_content_or_credentials(self):
        client = SimpleNamespace(request=Mock(return_value={"reviewId": "123_Ab-c", "isLiked": 0,
            "review": {"content": "private comment", "isLike": 1, "author": {"userVid": "456"}, "chapterUid": 3}}))
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            report = w.probe_review(client, "123_Ab-c")
        client.request.assert_called_once_with("/review/single", {"reviewId": "123_Ab-c"}, agent=True)
        self.assertEqual(report["state_fields"], {"response.isLiked": "0", "response.review.isLike": "1"})
        self.assertNotIn("private comment", output.getvalue())
        self.assertNotIn("456", output.getvalue())
        self.assertIn("isLike 已通过", output.getvalue())

    def test_missing_state_is_unknown_even_with_large_like_count(self):
        client = SimpleNamespace(request=Mock(return_value={"reviewId": "123_Ab-c", "review": {
            "content": "sample", "likesCount": 9999}}))
        with contextlib.redirect_stdout(io.StringIO()):
            report = w.probe_review(client, "123_Ab-c")
        self.assertEqual(report["state_fields"], {})

    def test_probe_rejects_different_review(self):
        client = SimpleNamespace(request=Mock(return_value={"reviewId": "321_Xy", "content": "sample"}))
        with self.assertRaisesRegex(w.ExportError, "不一致"):
            w.probe_review(client, "123_Ab-c")


class PresentationTests(unittest.TestCase):
    def setUp(self):
        self.book = w.demo_book(ZoneInfo("UTC"))

    def test_hidden_dates_are_absent_from_both_export_streams(self):
        self.book.show_dates = False
        content = "\n".join(text for _, text in w.entries(self.book))
        self.assertNotRegex(content, r"\d{4}-\d{2}-\d{2}")
        self.assertNotIn("划线时间", content)
        self.assertNotIn("日期未知", content)
        self.assertIn("我的想法", content)
        self.assertIn("读者观点", content)
        self.assertNotIn("bookId", content)
        self.assertFalse(any(kind == "meta" for kind, _ in w.entries(self.book)))
        for block in w.notion_blocks(self.book):
            text = w.block_signature(block)[1]
            self.assertNotRegex(text, r"\d{4}-\d{2}-\d{2}")

    def test_dates_follow_body_in_short_parentheses(self):
        self.book.show_dates = True
        content = list(w.entries(self.book))
        quote = next(text for kind, text in content if kind == "quote")
        thought = next(text for kind, text in content if kind == "thought")
        liked = next(text for kind, text in content if kind == "liked")
        for text in (quote, thought, liked):
            self.assertRegex(text, r"（\d{4}-\d{2}-\d{2}）$")
            self.assertNotRegex(text, r"\d{2}:\d{2}|\+\d{4}")
        self.assertEqual(thought.splitlines()[0], "我的想法")
        self.assertTrue(liked.startswith("读者观点\n"))
        self.assertNotIn("一位读者", liked)
        self.assertNotIn("128 赞", liked)
        unknown = self.book.chapters[-1].groups[0].highlight
        self.assertEqual(w.dated_text(unknown, True), unknown.text)

    def test_word_has_plain_quotes_and_clear_title_hierarchy(self):
        from docx import Document
        from docx.oxml.ns import qn
        self.book.show_dates = True
        with tempfile.TemporaryDirectory() as folder:
            document = Document(w.export_word(self.book, Path(folder)))
            self.assertEqual(document.paragraphs[0].text, self.book.title)
            self.assertEqual(document.paragraphs[0].style.name, "Heading 1")
            self.assertEqual(document.paragraphs[1].style.name, "Subtitle")
            self.assertEqual(document.paragraphs[1].text, self.book.author + " · 读书笔记")
            for name in ("Normal", "Heading 1", "Heading 2", "Quote", "Subtitle"):
                self.assertFalse(document.styles[name].font.italic)
            for paragraph in document.paragraphs:
                self.assertTrue(all(run.italic is not True for run in paragraph.runs))
            timestamp_runs = [run for p in document.paragraphs for run in p.runs
                              if w.re.fullmatch(r"（\d{4}-\d{2}-\d{2}）", run.text)]
            self.assertTrue(timestamp_runs)
            self.assertTrue(all(run.font.size.pt == 9 for run in timestamp_runs))
            self.assertTrue(all(p.style.name == "Heading 2" for p in document.paragraphs
                                if p.text in [c.title for c in self.book.chapters]))

    def test_notion_uses_plain_paragraphs_and_native_comment_plan(self):
        blocks, jobs = w.notion_plan(self.book)
        self.assertFalse(any(block["type"] in ("quote", "bulleted_list_item") for block in blocks))
        self.assertEqual(len(jobs), 2)
        self.assertEqual([job["rich_text"][0]["text"]["content"] for job in jobs], ["我的想法", "读者观点"])
        self.assertEqual(sum(block["type"] == "divider" for block in blocks), len(self.book.chapters) - 1)
        for job in jobs:
            self.assertEqual(blocks[job["anchor_index"]]["type"], "paragraph")

    def test_same_original_groups_my_thought_and_reader_without_duplicate_excerpt(self):
        book = w.build_notebook({"bookId": "b", "title": "书"}, [{"chapterUid": 1, "title": "章节"}], [],
            [{"reviewId": "m", "chapterUid": 1, "range": "1-10", "abstract": "同一段原文", "content": "我的批注"}],
            [{"reviewId": "r", "chapterUid": 1, "range": "1-10", "abstract": "同一段原文", "content": "他人观点",
              "author": {"name": "不应保留的名字"}, "likesCount": 9876}], ZoneInfo("UTC"))
        self.assertEqual(len(book.chapters[0].groups), 1)
        units = [unit for unit in w.display_units(book) if any(kind == "quote" for kind, _ in unit)]
        self.assertEqual(len(units), 1)
        self.assertEqual([kind for kind, _ in units[0]], ["quote", "thought", "liked"])
        text = "\n".join(value for _, value in w.entries(book))
        self.assertEqual(text.count("同一段原文"), 1)
        self.assertNotIn("不应保留的名字", text)
        self.assertNotIn("9876", text)
        for note in book.chapters[0].groups[0].thoughts:
            self.assertFalse(hasattr(note, "author"))
            self.assertFalse(hasattr(note, "likes"))
        blocks, jobs = w.notion_plan(book)
        original = [i for i, block in enumerate(blocks) if w.block_signature(block)[1] == "同一段原文"]
        self.assertEqual(len(original), 1)
        self.assertEqual([job["anchor_index"] for job in jobs], [original[0], original[0]])

    def test_different_originals_are_not_mixed(self):
        book = w.build_notebook({"bookId": "b", "title": "书"}, [{"chapterUid": 1, "title": "章"}], [],
            [{"reviewId": "m", "chapterUid": 1, "range": "1-10", "abstract": "原文甲", "content": "想法甲"}],
            [{"reviewId": "r", "chapterUid": 1, "range": "20-30", "abstract": "原文乙", "content": "观点乙"}], ZoneInfo("UTC"))
        units = [unit for unit in w.display_units(book) if any(kind == "quote" for kind, _ in unit)]
        self.assertEqual(len(units), 2)
        self.assertEqual([text for kind, text in units[0] if kind == "quote"], ["原文甲"])
        self.assertTrue(any("想法甲" in value for _, value in units[0]))
        self.assertFalse(any("观点乙" in value for _, value in units[0]))

    def test_many_scan_warnings_become_one_short_notice(self):
        self.book.warnings = ["读者观点仅覆盖可发现段落", "候选想法达到本次上限，部分内容可能未扫描。",
                              "发现 1000 条候选想法，仅核验前 200 条，可提高 --max-review-checks。",
                              "12 个章节、段落或想法未能读取，本次读者观点不完整。"]
        notices = [value for kind, value in w.entries(self.book) if kind == "notice"]
        self.assertEqual(notices, ["读者观点可能不完整。"])

    def test_word_connects_quote_and_comments_and_keeps_short_group_together(self):
        from docx import Document
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as folder:
            doc = Document(w.export_word(self.book, Path(folder)))
            quote_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith("阅读的价值"))
            paragraphs = doc.paragraphs[quote_index:quote_index + 3]
            for p in paragraphs:
                self.assertIsNotNone(p._p.find(".//" + qn("w:left")))
            self.assertTrue(paragraphs[0].paragraph_format.keep_with_next)
            self.assertTrue(paragraphs[1].paragraph_format.keep_with_next)
            self.assertFalse(paragraphs[2].paragraph_format.keep_with_next)

    def test_long_native_comments_respect_each_payload_limit(self):
        self.book.chapters[0].groups[0].thoughts[0].text = "长批注内容" * 20000
        blocks, jobs = w.notion_plan(self.book)
        self.assertGreater(len(jobs), 2)
        for job in jobs:
            self.assertLessEqual(len(job["rich_text"]), 21)
            self.assertLess(len(w.json.dumps(job).encode()), 300000)


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.tz = ZoneInfo("UTC")

    def test_nested_review_preserves_outer_like_status(self):
        row = w.unwrap_review({"isLiked": 0, "idx": 8, "review": {
            "reviewId": "x", "review": {"isLiked": 1, "content": "内容"}}})
        self.assertEqual(row["isLiked"], 0)
        self.assertEqual(row["content"], "内容")
        self.assertEqual(row["reviewId"], "x")
        self.assertEqual(row["idx"], 8)

    def test_cookie_header_normalization_and_redacted_errors(self):
        value, vid = w.parse_cookie('cookie: "wr_vid=123; wr_skey=example-only"')
        self.assertEqual(value, "wr_vid=123; wr_skey=example-only")
        self.assertEqual(vid, "123")
        for invalid in ("\x16", "wr_vid=123;\nwr_skey=example-only", "wr_vid=123; wr_skey="):
            with self.assertRaises(w.ExportError) as caught:
                w.parse_cookie(invalid)
            self.assertNotIn("example-only", str(caught.exception))

    def test_api_key_format_check(self):
        w.validate_api_key("wrk-example-test-only")
        for invalid in ("wr_vid=123", "wrk-", "wrk-example\n", "wrong"):
            with self.assertRaises(w.ExportError):
                w.validate_api_key(invalid)

    @patch("weread_notes.choice", return_value="1")
    @patch("weread_notes.secret", return_value="wrk-test-only")
    def test_key_switch_keeps_current_book_and_uses_gateway_for_personal_thoughts(self, secret, choice):
        client = w.Weread("wr_vid=123; wr_skey=example", interactive_auth=True)
        calls = []
        mark = {"bookmarkId": "h", "markText": "quote", "chapterUid": 1}
        def request(path, params=None, **kwargs):
            calls.append((path, params, kwargs))
            if path.endswith("/info"):
                return {"title": "书"}
            if path.endswith("chapterInfos"):
                return {"data": [{"bookId": "b", "updated": [{"chapterUid": 1, "title": "章节"}]}]}
            if path == "/api/user/notebook":
                return {"books": [{"bookId": "b", "noteCount": 1}]}
            if path == "/book/bookmarklist" and kwargs.get("agent"):
                return {"updated": [mark]}
            return {}
        pages_calls = []
        client.request = request
        def pages(path, *args, **kwargs):
            pages_calls.append((path, kwargs))
            return []
        client.pages = pages
        notebook = client.collect({"bookId": "b"}, ZoneInfo("UTC"), 0)
        self.assertEqual(notebook.chapters[0].groups[0].highlight.text, "quote")
        self.assertEqual(client.api_key, "wrk-test-only")
        self.assertEqual(pages_calls[0][0], "/review/list/mine")
        self.assertTrue(pages_calls[0][1]["agent"])
        gateway_calls = [call for call in calls if call[0] == "/book/bookmarklist" and call[2].get("agent")]
        self.assertEqual(gateway_calls[0][1]["bookId"], "b")

    @patch("weread_notes.choice", return_value="1")
    @patch("weread_notes.secret", return_value="wrk-test-only")
    def test_failed_key_switch_does_not_leave_changed_credentials(self, secret, choice):
        client = w.Weread("wr_vid=123; wr_skey=example", interactive_auth=True)
        def request(path, params=None, **kwargs):
            if kwargs.get("agent"):
                raise w.ExportError("测试：授权被拒绝")
            if path == "/api/user/notebook":
                return {"books": [{"bookId": "b", "noteCount": 1}]}
            return {}
        client.request = request
        with self.assertRaisesRegex(w.ExportError, "授权被拒绝"):
            client.highlights("b")
        self.assertEqual(client.api_key, "")

    @patch.dict(w.os.environ, {}, clear=True)
    @patch("weread_notes.read_clipboard", return_value="wr_vid=123; wr_skey=example-only")
    @patch("weread_notes.getpass.getpass")
    @patch("builtins.input", side_effect=["", ""])
    def test_clipboard_default_without_pasting_or_echoing(self, input_mock, getpass_mock, clipboard):
        output = io.StringIO()
        with patch("weread_notes.os.name", "nt"), contextlib.redirect_stdout(output):
            value = w.secret("Cookie：", "WEREAD_COOKIE", validator=w.parse_cookie)
        self.assertIn("wr_vid=123", value)
        self.assertNotIn("example-only", output.getvalue())
        self.assertIn("必需字段检查通过", output.getvalue())
        clipboard.assert_called_once()
        getpass_mock.assert_not_called()

    @patch.dict(w.os.environ, {}, clear=True)
    @patch("weread_notes.read_clipboard", side_effect=[w.ExportError("剪贴板暂时被占用"), "bad", "wr_vid=123; wr_skey=example-only"])
    @patch("builtins.input", side_effect=["1", "", "1", "", "1", ""])
    def test_clipboard_busy_and_invalid_cookie_can_retry(self, input_mock, clipboard):
        with patch("weread_notes.os.name", "nt"), contextlib.redirect_stdout(io.StringIO()):
            value = w.secret("Cookie：", "WEREAD_COOKIE", validator=w.parse_cookie)
        self.assertEqual(value, "wr_vid=123; wr_skey=example-only")
        self.assertEqual(clipboard.call_count, 3)

    @patch.dict(w.os.environ, {"WEREAD_COOKIE": "invalid"}, clear=True)
    @patch("weread_notes.read_clipboard")
    def test_invalid_environment_cookie_fails_without_reading_clipboard(self, clipboard):
        with self.assertRaisesRegex(w.ExportError, "环境变量 WEREAD_COOKIE 无效"):
            w.secret("Cookie：", "WEREAD_COOKIE", validator=w.parse_cookie)
        clipboard.assert_not_called()

    def test_windows_clipboard_unlocks_and_closes_without_accessing_real_clipboard(self):
        import ctypes
        payload = "wr_vid=123; wr_skey=example-only\0".encode("utf-16-le")
        buffer = ctypes.create_string_buffer(payload, len(payload))
        user = SimpleNamespace(OpenClipboard=Mock(return_value=True), CloseClipboard=Mock(),
                               GetClipboardData=Mock(return_value=42))
        kernel = SimpleNamespace(GlobalSize=Mock(return_value=len(payload)),
                                 GlobalLock=Mock(return_value=ctypes.addressof(buffer)), GlobalUnlock=Mock())
        with patch("weread_notes.os.name", "nt"), patch.object(ctypes, "WinDLL", create=True,
                                                             side_effect=[user, kernel]):
            self.assertEqual(w.read_clipboard(), "wr_vid=123; wr_skey=example-only")
        kernel.GlobalUnlock.assert_called_once_with(42)
        user.CloseClipboard.assert_called_once()

    def test_chapter_response_variants(self):
        chapter = {"chapterUid": 7, "title": "章节"}
        for data in ({"chapters": [chapter]}, {"chapterInfos": [chapter]},
                     {"data": [{"bookId": "b", "updated": [chapter]}]},
                     {"data": [{"bookId": "other", "updated": []},
                               {"bookId": "b", "chapterInfos": [chapter]}]}):
            self.assertEqual(w.extract_chapters(data, "b"), [chapter])
        with self.assertRaises(w.ExportError):
            w.extract_chapters({"data": []}, "b")

    def test_highlight_response_and_explicit_empty_variants(self):
        record = {"bookmarkId": "a", "markText": "private test content"}
        for data in ({"updated": [record]}, {"data": {"updated": [record]}}):
            self.assertEqual(w.extract_highlights(data)[0], [record])
        self.assertEqual(w.extract_highlights({"totalCount": 0})[0], [])
        for data in ({}, {"synckey": 10}, {"updated": None}, {"updated": [None]}):
            self.assertIsNone(w.extract_highlights(data))

    def test_missing_updated_only_treated_empty_when_overview_confirms_zero(self):
        client = w.Weread("wr_vid=123; wr_skey=placeholder")
        query_calls = []
        def request(path, params=None, **kwargs):
            query_calls.append((path, params))
            if path.endswith("bookmarklist"):
                return {"synckey": 10}
            return {"books": [{"book": {"bookId": "b"}, "noteCount": 0}]}
        client.request = request
        self.assertEqual(client.highlights("b")[0], [])
        self.assertEqual(query_calls[0][1]["synckey"], 0)

    def test_missing_updated_with_positive_count_fails(self):
        client = w.Weread("wr_vid=123; wr_skey=placeholder")
        client.request = lambda path, *a, **kw: ({"synckey": 10} if path.endswith("bookmarklist") else
                                                {"books": [{"bookId": "b", "noteCount": 3}]})
        with self.assertRaisesRegex(w.ExportError, "该书有划线"):
            client.highlights("b")

    def test_book_absent_from_overview_does_not_imply_empty(self):
        client = w.Weread("wr_vid=123; wr_skey=placeholder")
        client.request = lambda path, *a, **kw: ({"synckey": 10} if path.endswith("bookmarklist") else
                                                {"books": []})
        with self.assertRaisesRegex(w.ExportError, "无法确认"):
            client.highlights("b")

    def test_response_shape_contains_no_note_content_or_cookie_values(self):
        schema = w.response_shape({"book": {"title": "private book title"},
                                   "updated": [{"markText": "private note text"}], "token": "secret-token"})
        self.assertIn("updated:array(1)", schema)
        self.assertNotIn("private", schema)
        self.assertNotIn("secret-token", schema)

    def test_chapter_order_and_unknown_fallback(self):
        book = w.demo_book(self.tz)
        self.assertEqual([c.uid for c in book.chapters], ["19", "3", "__unknown__"])
        group = book.chapters[0].groups[0]
        self.assertEqual([n.kind for n in group.thoughts], ["thought", "liked"])
        self.assertEqual(book.chapters[-1].title, "未分类章节")

    def test_duplicate_highlights_do_not_drop_distinct_ids(self):
        rows = [{"bookmarkId": "a", "text": "same"}, {"bookmarkId": "b", "text": "same"},
                {"bookmarkId": "a", "text": "same"}]
        self.assertEqual(len(w.unique_rows(rows, "bookmarkId")), 2)

    def test_unknown_chapters_do_not_cross_link_identical_quotes(self):
        book = w.build_notebook({"bookId": "b", "title": "书"}, [],
            [{"bookmarkId": "a", "chapterUid": 100, "markText": "same", "range": "1-2"}],
            [{"reviewId": "r", "chapterUid": 200, "abstract": "same", "content": "想法", "range": "1-2"}],
            [], self.tz)
        self.assertEqual(len(book.chapters[0].groups), 2)

    def test_dates_seconds_milliseconds_and_invalid(self):
        self.assertEqual(w.date_text(1700000000, self.tz), w.date_text(1700000000000, self.tz))
        self.assertEqual(w.date_text("bad", self.tz), "日期未知")
        self.assertEqual(w.date_text(0, self.tz), "日期未知")

    def test_notion_id_ignores_view_id(self):
        bid, vid = "a" * 32, "b" * 32
        self.assertEqual(w.notion_id(f"https://www.notion.so/Title-{bid}?v={vid}"), str(w.uuid.UUID(bid)))
        with self.assertRaises(w.ExportError):
            w.notion_id("not-an-id")

    def test_unicode_chunking_preserves_every_character(self):
        value = "长文🙂" * 3000
        parts = w.text_chunks(value)
        self.assertEqual("".join(parts), value)
        self.assertTrue(all(len(p.encode("utf-16-le")) // 2 <= 1800 for p in parts))

    def test_public_filter_excludes_unliked_own_and_unattributed(self):
        client = w.Weread("wr_vid=123; wr_skey=placeholder")
        client.request = lambda path, *a, **kw: (
            {"title": "测试书"} if path.endswith("/info") else
            {"data": [{"bookId": "b", "updated": [{"chapterUid": 1, "title": "第一章"}]}]}
            if path.endswith("chapterInfos") else {"updated": [], "chapters": []})
        public = [
            {"review": {"reviewId": "liked", "isLiked": 1, "userVid": 456, "likesCount": 99, "chapterUid": 1, "content": "保留"}},
            {"review": {"reviewId": "unliked", "isLiked": 0, "userVid": 456, "likesCount": 9999, "chapterUid": 1, "content": "未点赞"}},
            {"review": {"reviewId": "own", "isLiked": 1, "userVid": 123, "chapterUid": 1, "content": "自己"}},
            {"review": {"reviewId": "missing", "isLiked": 1, "chapterUid": 1, "content": "无作者"}},
            {"review": {"reviewId": "no-status", "userVid": 456, "chapterUid": 1, "content": "无状态"}},
            {"review": {"reviewId": "low", "isLiked": 1, "userVid": 456, "likesCount": 2, "chapterUid": 1, "content": "门槛以下"}},
        ]
        client.pages = lambda path, query, *a, **kw: [] if query.get("mine") == 1 else public
        book = client.collect({"bookId": "b"}, self.tz, 10)
        notes = [n for c in book.chapters for g in c.groups for n in g.thoughts]
        self.assertEqual([n.source_id for n in notes], ["liked"])
        self.assertTrue(any("isLiked" in message for message in book.warnings))
        self.assertTrue(any("作者身份" in message for message in book.warnings))

    def test_real_shelf_not_notebooks(self):
        client = w.Weread("wr_vid=123; wr_skey=placeholder")
        paths = []
        def request(path, *args, **kwargs):
            paths.append(path)
            return {"books": [{"bookId": "b", "title": "尚未记笔记的书"}]}
        client.request = request
        self.assertEqual(client.shelf()[0]["bookId"], "b")
        self.assertEqual(paths, ["/web/shelf/sync"])

    def test_pagination_and_stalled_cursor(self):
        client = w.Weread("wr_vid=123; wr_skey=placeholder")
        responses = iter([
            {"reviews": [{"idx": 1}], "hasMore": 1, "synckey": 2},
            {"reviews": [{"idx": 2}], "hasMore": 0},
        ])
        queries = []
        def request(path, query, **kwargs):
            queries.append(dict(query))
            return next(responses)
        client.request = request
        self.assertEqual(len(client.pages("/x", {"synckey": 0}, "reviews")), 2)
        self.assertEqual(queries[1]["maxIdx"], 1)
        self.assertEqual(queries[1]["synckey"], 2)
        client.request = lambda *a, **kw: {"reviews": [{"idx": 1}], "hasMore": 1, "synckey": 0}
        with self.assertRaises(w.ExportError):
            client.pages("/x", {"synckey": 0, "maxIdx": 1}, "reviews")

    def test_word_headings_and_non_overwrite(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            book = w.demo_book(self.tz)
            first = w.export_word(book, Path(folder))
            second = w.export_word(book, Path(folder))
            self.assertNotEqual(first, second)
            doc = Document(first)
            self.assertEqual(doc.paragraphs[0].style.name, "Heading 1")
            headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
            self.assertEqual(headings, [c.title for c in book.chapters])
            self.assertTrue(any(p.style.name == "Quote" for p in doc.paragraphs))
            self.assertIn("我的想法", "\n".join(p.text for p in doc.paragraphs))


class FakeNotion:
    """模拟远端写入成功但客户端超时，重跑时应检测已存在块。"""
    def __init__(self):
        self.created, self.blocks_written, self.appends, self.fail_after_write = 0, [], [], False
        self.comment_rows, self.fail_comment_after_write = [], False
        self.comments = SimpleNamespace(list=self.list_comments, create=self.create_comment)
        self.title_property = "书名"
        self.pages = SimpleNamespace(retrieve=lambda **kw: {}, create=self.create)
        self.databases = SimpleNamespace(retrieve=lambda **kw: {"data_sources": [{"id": "c" * 32}]})
        self.data_sources = SimpleNamespace(retrieve=lambda **kw: {"properties": {self.title_property: {"type": "title"}}})
        self.blocks = SimpleNamespace(children=SimpleNamespace(list=self.list_blocks, append=self.append))

    def create(self, **kwargs):
        self.created += 1
        self.created_kwargs = kwargs
        return {"id": "d" * 32, "url": "https://www.notion.so/" + "d" * 32}

    def list_blocks(self, **kwargs):
        offset = int(kwargs.get("start_cursor", 0))
        rows = self.blocks_written[offset:offset + 100]
        more = offset + len(rows) < len(self.blocks_written)
        return {"results": rows, "has_more": more, "next_cursor": str(offset + len(rows)) if more else None}

    def append(self, **kwargs):
        self.appends.append(kwargs["children"])
        for block in kwargs["children"]:
            self.blocks_written.append({**w.json.loads(w.json.dumps(block)), "id": f"block-{len(self.blocks_written)}"})
        if self.fail_after_write:
            self.fail_after_write = False
            raise TimeoutError("simulated after server committed")
        return {}

    def list_comments(self, **kwargs):
        rows = [row for row in self.comment_rows if row["parent"]["block_id"] == kwargs["block_id"]]
        offset = int(kwargs.get("start_cursor", 0))
        page = rows[offset:offset + 100]
        more = offset + len(page) < len(rows)
        return {"results": page, "has_more": more, "next_cursor": str(offset + len(page)) if more else None}

    def create_comment(self, **kwargs):
        row = {**w.json.loads(w.json.dumps(kwargs)), "id": f"comment-{len(self.comment_rows)}"}
        self.comment_rows.append(row)
        if self.fail_comment_after_write:
            self.fail_comment_after_write = False
            raise TimeoutError("simulated after comment committed")
        return row


class NotionTests(unittest.TestCase):
    def setUp(self):
        self.book = w.demo_book(ZoneInfo("UTC"))
        self.fake = FakeNotion()

    @patch("weread_notes.time.sleep")
    def test_database_schema_discovery_and_idempotent_rerun(self, _):
        with tempfile.TemporaryDirectory() as folder:
            args = (self.book, "not-a-real-token", "a" * 32, "database", Path(folder))
            first = w.export_notion(*args, client=self.fake)
            count = len(self.fake.blocks_written)
            self.assertIn("书名", self.fake.created_kwargs["properties"])
            self.assertEqual(self.fake.created_kwargs["parent"]["type"], "data_source_id")
            self.assertEqual(w.export_notion(*args, client=self.fake), first)
            self.assertEqual(len(self.fake.blocks_written), count)
            self.assertEqual(self.fake.created, 1)

    @patch("weread_notes.time.sleep")
    def test_resume_after_server_commit_and_client_timeout(self, _):
        with tempfile.TemporaryDirectory() as folder:
            args = (self.book, "not-a-real-token", "a" * 32, "page", Path(folder))
            self.fake.fail_after_write = True
            with self.assertRaises(w.ExportError):
                w.export_notion(*args, client=self.fake)
            expected = len(w.notion_blocks(self.book))
            self.assertEqual(len(self.fake.blocks_written), expected)
            w.export_notion(*args, client=self.fake)
            self.assertEqual(len(self.fake.blocks_written), expected)
            self.assertEqual(len(self.fake.appends), 1)

    @patch("weread_notes.time.sleep")
    def test_long_payload_chunk_limits(self, _):
        # 触发 > 100 块、> 350KB 的分批；含中文和 emoji。
        self.book.chapters[0].groups[0].highlight.text = "很长的摘录🙂" * 25000
        with tempfile.TemporaryDirectory() as folder:
            w.export_notion(self.book, "fake", "a" * 32, "page", Path(folder), client=self.fake)
            for batch in self.fake.appends:
                self.assertLessEqual(len(batch), 100)
                self.assertLess(len(w.json.dumps({"children": batch}).encode()), 500000)
            self.assertEqual(len(self.fake.blocks_written), len(w.notion_blocks(self.book)))

    @patch("weread_notes.time.sleep")
    def test_refuses_resume_into_edited_page(self, _):
        with tempfile.TemporaryDirectory() as folder:
            args = (self.book, "fake", "a" * 32, "page", Path(folder))
            w.export_notion(*args, client=self.fake)
            self.fake.blocks_written[0]["paragraph"]["rich_text"][0]["text"]["content"] = "用户编辑过"
            with self.assertRaisesRegex(w.ExportError, "已被修改"):
                w.export_notion(*args, client=self.fake)

    @patch("weread_notes.time.sleep")
    def test_rerun_preserves_user_edited_native_comment(self, _):
        with tempfile.TemporaryDirectory() as folder:
            args = (self.book, "fake", "a" * 32, "page", Path(folder))
            w.export_notion(*args, client=self.fake)
            comment = self.fake.comment_rows[0]
            comment["rich_text"][0]["text"]["content"] = "手工编辑了批注"
            count = len(self.fake.comment_rows)
            w.export_notion(*args, client=self.fake)
            self.assertEqual(len(self.fake.comment_rows), count)
            self.assertEqual(comment["rich_text"][0]["text"]["content"], "手工编辑了批注")

    @patch("weread_notes.time.sleep")
    def test_native_comments_are_anchored_to_exact_source_blocks(self, _):
        with tempfile.TemporaryDirectory() as folder:
            w.export_notion(self.book, "fake", "a" * 32, "page", Path(folder), client=self.fake)
            blocks, jobs = w.notion_plan(self.book)
            self.assertEqual(len(self.fake.comment_rows), len(jobs))
            for comment, job in zip(self.fake.comment_rows, jobs):
                self.assertEqual(comment["parent"]["block_id"], self.fake.blocks_written[job["anchor_index"]]["id"])
                self.assertEqual(comment["rich_text"], job["rich_text"])
            self.assertFalse(any(block["type"] == "quote" for block in self.fake.blocks_written))

    @patch("weread_notes.time.sleep")
    def test_comment_timeout_after_commit_does_not_duplicate_on_resume(self, _):
        with tempfile.TemporaryDirectory() as folder:
            args = (self.book, "fake", "a" * 32, "page", Path(folder))
            self.fake.fail_comment_after_write = True
            with self.assertRaisesRegex(w.ExportError, "原生评论未完成"):
                w.export_notion(*args, client=self.fake)
            self.assertEqual(len(self.fake.comment_rows), 1)
            count = len(self.fake.blocks_written)
            w.export_notion(*args, client=self.fake)
            self.assertEqual(len(self.fake.comment_rows), len(w.notion_plan(self.book)[1]))
            self.assertEqual(len(self.fake.blocks_written), count)

    @patch("weread_notes.time.sleep")
    def test_specific_page_append_keeps_existing_content_and_does_not_create_page(self, _):
        original = {"id": "original-block", "type": "paragraph", "paragraph": {
            "rich_text": [{"text": {"content": "原有内容"}}]}}
        self.fake.blocks_written.append(original)
        with tempfile.TemporaryDirectory() as folder:
            args = (self.book, "fake", "a" * 32, "existing_page", Path(folder))
            w.export_notion(*args, client=self.fake)
            self.assertEqual(self.fake.created, 0)
            self.assertEqual(self.fake.blocks_written[0], original)
            self.assertEqual(self.fake.blocks_written[1]["type"], "divider")
            block_count, comment_count = len(self.fake.blocks_written), len(self.fake.comment_rows)
            w.export_notion(*args, client=self.fake)
            self.assertEqual(len(self.fake.blocks_written), block_count)
            self.assertEqual(len(self.fake.comment_rows), comment_count)


class DestinationTests(unittest.TestCase):
    @patch.dict(w.os.environ, {}, clear=True)
    @patch("weread_notes.choice", side_effect=["1", "2"])
    @patch("weread_notes.ask", return_value="a" * 32)
    def test_new_note_under_parent(self, ask, choice):
        self.assertEqual(w.ask_notion_destination(), ("page", "a" * 32))
        self.assertIn("新建笔记", choice.call_args_list[0].args[0])

    @patch.dict(w.os.environ, {}, clear=True)
    @patch("weread_notes.choice", side_effect=["2", "1"])
    @patch("weread_notes.ask", return_value="a" * 32)
    def test_specific_existing_note_is_append_target(self, ask, choice):
        self.assertEqual(w.ask_notion_destination(), ("existing_page", "a" * 32))
        self.assertIn("指定位置", choice.call_args_list[0].args[0])


if __name__ == "__main__":
    # 测试只显示结果，不显示模拟的终端进度。
    with contextlib.redirect_stdout(io.StringIO()):
        unittest.main(verbosity=2)
