"""只添加的增量同步：章节内精确去重，保留手工修改，不根据相似度删除。"""
from __future__ import annotations
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import os
import shutil
import tempfile
import unicodedata
import weread_notes as core


def normalized(text):
    # 统一空白和 Unicode 表示，不删除标点、句子或原文中的日期。
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFC', text or '')).strip()


def identifier(*values):
    return hashlib.sha256(json.dumps(values, ensure_ascii=False).encode()).hexdigest()


class SyncError(core.ExportError):
    def __init__(self, code, message, choices=None):
        super().__init__(message)
        self.code, self.choices = code, choices or []


def records_for(book):
    """稳定 ID 不依赖显示日期；原文和评论分开记账，避免新评论被旧原文吞掉。"""
    records, seen = [], set()
    for chapter in book.chapters:
        notes = [n for group in chapter.groups for n in group.thoughts]
        highlights = [g.highlight for g in chapter.groups if g.highlight]
        sample = core.Notebook(book.book_id, book.title, book.author, [chapter], [], book.show_dates)
        for unit in core.display_units(sample):
            source_key, source_variants = None, []
            for kind, text in unit:
                if kind == 'quote':
                    note = next((n for n in highlights if core.dated_text(n, book.show_dates) == text), None)
                    raw = note.text if note else text
                    source_key = identifier(book.book_id, chapter.uid, 'source', normalized(raw))
                    source_variants = list(dict.fromkeys([raw, text] + ([core.dated_text(note, True)] if note else [])))
                    record = {'key': source_key, 'kind': 'source', 'text': text, 'variants': source_variants}
                elif kind in ('thought', 'liked'):
                    label, _, body = text.partition('\n')
                    note = next((n for n in notes if n.kind == kind and core.dated_text(n, book.show_dates) == body
                                 and (not n.quote or normalized(n.quote) in {normalized(s) for s in source_variants})), None)
                    raw = note.text if note else body
                    source_id = note.source_id if note and note.source_id not in ('', 'None') else None
                    key = identifier(book.book_id, chapter.uid, kind, source_id or [source_key, normalized(raw)])
                    variants = [label + '\n' + raw, text]
                    if note:
                        variants.append(label + '\n' + core.dated_text(note, True))
                    record = {'key': key, 'kind': kind, 'text': text, 'variants': list(dict.fromkeys(variants)),
                              'source_key': source_key, 'source_variants': source_variants,
                              'label': label, 'body': body}
                else:
                    continue
                record.update(chapter_uid=chapter.uid, chapter_title=chapter.title, chapter_order=chapter.order)
                if record['key'] not in seen:
                    records.append(record)
                    seen.add(record['key'])
    return records


def comment_parts(record):
    chunks, jobs = core.text_chunks(record['body']), []
    for offset in range(0, len(chunks), 20):
        rich = [{'type': 'text', 'text': {'content': record['label'] + ('（续）' if offset else '')},
                 'annotations': {'bold': True}}]
        rich.extend({'type': 'text', 'text': {'content': ('\n' if i == 0 else '') + part}}
                    for i, part in enumerate(chunks[offset:offset + 20]))
        jobs.append({'part': offset // 20, 'rich_text': rich})
    return jobs


def notion_delta(plan, snapshot, receipt=None):
    """输入实际读取的远端快照，输出仅新增的正文/评论；不执行远端写入。

    snapshot: {book_id,page_id,chapters:[{title,paragraphs:[{text,id?}]}],
               comments:[{id?,chapter_title,selection,text}],comments_complete:true}
    receipt: {book_id,page_id,completed:[{key,remote_id?}]}，写入确认后才记录。
    """
    if not isinstance(plan.get('records'), list):
        raise SyncError('old_plan', '旧版导出计划不含增量身份，请重新抓取生成新版计划。')
    if snapshot.get('book_id') != plan['book_id'] or not snapshot.get('page_id'):
        raise SyncError('wrong_destination', '快照必须对应当前书籍及具体目标页面。')
    if not snapshot.get('comments_complete') and any(r['kind'] != 'source' for r in plan['records']):
        raise SyncError('comments_not_read', '请先完整读取目标页的评论，再进行增量对比，避免重复评论。')
    receipt = receipt or {}
    if receipt and (receipt.get('book_id') != plan['book_id'] or receipt.get('page_id') != snapshot['page_id']):
        raise SyncError('wrong_receipt', '同步记录属于其他书籍或页面，不能混用。')
    completed = {r['key']: r for r in receipt.get('completed', [])}
    completed_parts = {(r['key'], r['part']) for r in receipt.get('comment_parts', [])}
    chapters = {}
    for chapter in snapshot.get('chapters', []):
        name = normalized(chapter['title'])
        if name in chapters:
            raise SyncError('ambiguous_chapter', '目标页有同名章节，需要先在对话中确认要写入哪个章节。')
        chapters[name] = chapter
    result = {'schema_version': 1, 'book_id': plan['book_id'], 'page_id': snapshot['page_id'],
              'plan_fingerprint': plan['fingerprint'], 'chapters': [], 'comments': [],
              'retained': 0, 'matched': [], 'unresolved': []}
    additions, sources = {}, {}
    all_paragraphs = [p for c in chapters.values() for p in c.get('paragraphs', [])]
    for record in plan['records']:
        name = normalized(record['chapter_title'])
        chapter = chapters.get(name)
        variants = {normalized(v) for v in record['variants']}
        previous = completed.get(record['key'])
        if record['kind'] == 'source':
            matches = [p for p in (chapter or {}).get('paragraphs', []) if normalized(p['text']) in variants]
            if previous and previous.get('remote_id'):
                by_id = [p for p in all_paragraphs if p.get('id') == previous['remote_id']]
                matches = by_id or matches
            if matches or previous:
                result['retained'] += 1
                result['matched'].append({'key': record['key'], 'remote_id': matches[0].get('id') if len(matches) == 1 else None})
                # 旧回执证明导入过，但原文被手动改写/删除且无可定位 ID 时，不重建原文。
                sources[record['key']] = matches[0] if len(matches) == 1 else None
                continue
            group = additions.setdefault(name, {'uid': record['chapter_uid'], 'title': record['chapter_title'],
                                                'order': record['chapter_order'],
                                                'new_chapter': chapter is None, 'paragraphs': []})
            group['paragraphs'].append({'key': record['key'], 'text': record['text']})
            sources[record['key']] = {'text': record['text'], 'new': True}
        else:
            matches = [c for c in snapshot.get('comments', [])
                       if normalized(c.get('chapter_title')) == name and normalized(c.get('text')) in variants
                       and ((not record['source_key'] and not c.get('selection')) or
                            normalized(c.get('selection')) in {normalized(v) for v in record['source_variants']})]
            if previous or matches:
                result['retained'] += 1
                result['matched'].append({'key': record['key'], 'remote_id': matches[0].get('id') if matches else previous.get('remote_id')})
                continue
            anchor = sources.get(record['source_key']) if record['source_key'] else {'text': record['chapter_title']}
            if not anchor:
                result['unresolved'].append({'key': record['key'], 'reason': '对应原文已修改、删除或有多个匹配，请先定位原文再关联新评论。'})
                continue
            jobs = [job for job in comment_parts(record) if (record['key'], job['part']) not in completed_parts]
            if not jobs:
                result['retained'] += 1
                result['matched'].append({'key': record['key'], 'remote_id': None})
                continue
            result['comments'].append({'key': record['key'], 'chapter_title': record['chapter_title'],
                                       'selection': anchor['text'], 'anchor_block_id': anchor.get('id'),
                                       'source_key': record['source_key'], 'jobs': jobs})
    result['chapters'] = list(additions.values())
    for chapter in result['chapters']:
        next_titles = [r['chapter_title'] for r in plan['records'] if r['chapter_order'] > chapter['order']
                       and normalized(r['chapter_title']) in chapters]
        chapter['next_existing_chapter'] = next_titles[0] if next_titles else None
    result['added_sources'] = sum(len(c['paragraphs']) for c in result['chapters'])
    result['added_comments'] = len(result['comments'])
    return result


def export_word_incremental(book, output, target=None, fresh=False):
    """Word 通过隐藏书签记录身份；复制新增段落，既有 XML/人工内容不重建。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    index = output / ('word-sync-' + identifier(book.book_id)[:20] + '.json')
    records = records_for(book)
    state = json.loads(index.read_text(encoding='utf-8')) if index.exists() else {}
    if target:
        path = Path(target).resolve()
        if not path.is_file() or path.suffix.lower() != '.docx':
            raise SyncError('word_target_missing', '请提供现有 Word 笔记文件。')
    elif not fresh and state.get('file'):
        path = output / Path(state['file']).name
        if not path.is_file():
            raise SyncError('word_target_missing', '上次同步的 Word 文件已移动，请在对话中指定新位置，或明确选择新建笔记。')
    else:
        pattern = re.compile(r'^《' + re.escape(core.safe_filename(book.title)) + r'》- 读书笔记(?: \(\d+\))?\.docx$')
        candidates = [] if fresh else [p for p in output.iterdir() if pattern.fullmatch(p.name)]
        if len(candidates) > 1:
            raise SyncError('word_target_required', '发现多份同名笔记，请在对话中选定要更新的文件。', [str(p.resolve()) for p in candidates])
        path = candidates[0] if candidates else None
    new_document = path is None
    if new_document:
        path = core.export_word(book, output)
    doc = Document(path)
    bookmark_ids = [int(e.get(qn('w:id'))) for e in doc.element.iter(qn('w:bookmarkStart'))
                    if (e.get(qn('w:id')) or '').isdigit()]
    next_id = max(bookmark_ids, default=0) + 1

    def marked(key):
        name = '_wr_' + key[:32]
        for paragraph in doc.paragraphs:
            if any(e.get(qn('w:name')) == name for e in paragraph._p.iter(qn('w:bookmarkStart'))):
                return paragraph

    def mark(paragraph, key):
        nonlocal next_id
        start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
        start.set(qn('w:id'), str(next_id))
        start.set(qn('w:name'), '_wr_' + key[:32])
        end.set(qn('w:id'), str(next_id))
        paragraph._p.append(start)
        paragraph._p.append(end)
        next_id += 1

    def in_chapter(title):
        paragraphs, active, heads = [], False, 0
        for p in doc.paragraphs:
            if p.style.name == 'Heading 2':
                active = normalized(p.text) == normalized(title)
                heads += active
            elif active:
                paragraphs.append(p)
        if heads > 1:
            raise SyncError('ambiguous_chapter', 'Word 中存在同名章节，不能自动决定增量写入位置。')
        return paragraphs

    def matching(record):
        variants = {normalized(v) for v in record['variants']}
        candidates = in_chapter(record['chapter_title'])
        if record['kind'] != 'source' and record.get('source_key'):
            anchor = marked(record['source_key'])
            following, active = [], False
            for p in candidates:
                if anchor is not None and p._p is anchor._p:
                    active = True
                    continue
                if active and p.style.name == 'Quote':
                    break
                if active:
                    following.append(p)
            candidates = following
        matches = [p for p in candidates if normalized(p.text) in variants]
        if len(matches) > 1:
            raise SyncError('ambiguous_source', '同一章节有多处完全相同的原文或批注，请先在对话中确认位置。')
        return matches[0] if matches else None

    with tempfile.TemporaryDirectory() as folder:
        template = Document(core.export_word(book, Path(folder)))
        def template_element(record):
            active = False
            for p in template.paragraphs:
                if p.style.name == 'Heading 2':
                    active = p.text == record['chapter_title']
                elif active and p.text == record['text']:
                    return deepcopy(p._p)
            raise SyncError('template_missing', '新增段落排版失败，未保存已有文档。')

        def append_to_chapter(title, element):
            active, found = False, False
            for p in doc.paragraphs:
                if p.style.name != 'Heading 2':
                    continue
                if active:
                    p._p.addprevious(element)
                    return
                if normalized(p.text) == normalized(title):
                    active, found = True, True
            if not found:
                heading = next(p for p in template.paragraphs if p.style.name == 'Heading 2' and p.text == title)
                doc.element.body.insert(len(doc.element.body) - 1, deepcopy(heading._p))
            doc.element.body.insert(len(doc.element.body) - 1, element)

        retained, added, changed = 0, 0, False
        for record in records:
            if marked(record['key']) is not None:
                retained += 1
                continue
            existing = matching(record)
            if existing is not None:
                mark(existing, record['key'])
                retained += 1
                changed = True
                continue
            element = template_element(record)
            anchor = marked(record.get('source_key', '')) if record.get('source_key') else None
            if record['kind'] != 'source' and anchor is not None:
                sibling = anchor._p
                while sibling.getnext() is not None:
                    node = sibling.getnext()
                    plain = ''.join(node.itertext()) if node.tag == qn('w:p') else ''
                    if not plain.startswith(('我的想法', '读者观点')):
                        break
                    sibling = node
                sibling.addnext(element)
            else:
                append_to_chapter(record['chapter_title'], element)
            from docx.text.paragraph import Paragraph
            mark(Paragraph(element, doc._body), record['key'])
            added += 1
            changed = True
        if changed:
            descriptor, temporary = tempfile.mkstemp(suffix='.docx', dir=path.parent)
            os.close(descriptor)
            try:
                doc.save(temporary)
                if not new_document:
                    shutil.copyfile(path, str(path) + '.before-sync.bak')
                os.replace(temporary, path)
            finally:
                Path(temporary).unlink(missing_ok=True)
    # 外部指定路径不放入相对文件索引，之后可再显式指定；不误指向输出目录内的同名文件。
    if path.resolve().parent == output.resolve():
        core.atomic_json(index, {'book_id': book.book_id, 'file': path.name})
    return {'path': str(path.resolve()), 'added': len(records) if new_document else added,
            'retained': 0 if new_document else retained, 'incremental': not new_document}
