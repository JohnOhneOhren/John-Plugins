#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""微信读书交互式笔记导出器，Python 3.10+。

运行：python weread_notes.py；离线演示：python weread_notes.py --demo
凭证从环境变量、用户主动选择的剪贴板或隐藏输入读取，不写入文件。
所有微信接口封装在 Weread 中，便于接口变动时维护。
注意：listType=11 不代表“已点赞”。官方单条接口的 isLike 已经由点赞/取消对照验证。
"""
from __future__ import annotations

import argparse
import getpass
import hashlib
import html
import json
import os
import re
import sys
import time
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from http.cookiejar import Cookie, CookieJar, DefaultCookiePolicy
from http.cookies import SimpleCookie
from pathlib import Path
from typing import Any, Callable
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qs, urlencode, urlsplit
from urllib.request import HTTPCookieProcessor, HTTPRedirectHandler, Request, build_opener
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError


class ExportError(Exception):
    """可安全展示给用户的错误（不要包含服务端正文和凭证）。"""


def parse_cookie(value: str) -> tuple[str, str]:
    """规范化请求头粘贴内容，只报告字段名称，不回显凭证值。"""
    value = re.sub(r"^cookie\s*:\s*", "", value.strip(), flags=re.I)
    if len(value) >= 2 and value[0] == value[-1] and value[0] in ('"', "'"):
        value = value[1:-1].strip()
    if any(ord(c) < 32 for c in value):
        raise ExportError("Cookie 中含换行或控制字符。请复制单行 Cookie 完整值，选择从剪贴板读取。")
    parsed = SimpleCookie()
    try:
        parsed.load(value)
    except Exception:
        raise ExportError("Cookie 格式不正确，请复制请求标头 Cookie 的完整值。") from None
    missing = [key for key in ("wr_vid", "wr_skey") if key not in parsed or not parsed[key].value]
    if missing:
        raise ExportError("Cookie 缺少 " + "、".join(missing) + "。请复制完整请求标头值，不要复制单个条目。")
    return value, parsed["wr_vid"].value


def validate_api_key(value: str):
    if not value.startswith("wrk-") or len(value) <= 4 or any(c.isspace() or ord(c) < 32 for c in value):
        raise ExportError("请输入微信读书官方页面提供的完整 wrk- 开头 API Key，不是 Cookie。")


def review_id_from_reference(value: str) -> str:
    """只从明确的 reviewId 提取 ID，不把书籍/阅读器 hash 当作想法 ID。"""
    value = value.strip()
    if re.fullmatch(r"[0-9]+_[A-Za-z0-9_-]+", value):
        return value
    if "://" in value:
        parsed = urlsplit(value)
        if parsed.hostname not in ("weread.qq.com", "i.weread.qq.com"):
            raise ExportError("请提供微信读书原始想法链接或 reviewId，不支持其它网站链接。")
        query = parse_qs(parsed.query)
        for key in ("reviewId", "reviewid"):
            candidate = query.get(key, [""])[0]
            if re.fullmatch(r"[0-9]+_[A-Za-z0-9_-]+", candidate):
                return candidate
    raise ExportError("该链接未包含可识别的 reviewId，不能把阅读器地址当作想法 ID。"
                      "请复制想法本身的详情链接；短链接需要进一步解析。")


def probe_review(client: "Weread", reference: str) -> dict:
    """验证单条已知想法的状态字段，不修改点赞状态。

    输出仅含少量字段存在性与 0/1，不含正文、作者姓名、凭证或完整响应。
    """
    identifier = review_id_from_reference(reference)
    data = client.request("/review/single", {"reviewId": identifier}, agent=True)
    row = unwrap_review(data)
    if row.get("reviewId") and str(row["reviewId"]) != identifier:
        raise ExportError("单条详情返回的 reviewId 与请求不一致，不能用于验证。")
    if not (row.get("reviewId") or row.get("content") or row.get("htmlContent")):
        raise ExportError("单条想法详情没有返回可识别的想法数据。结构：" + response_shape(data))
    # 分层展示，避免把外层与内层的不同状态悄悄合并，或根据赞数猜测是否点赞。
    states = {}
    layer, path = data, "response"
    for _ in range(6):
        for key in ("isLiked", "isLike"):
            if key in layer:
                value = layer[key]
                states[path + "." + key] = str(value) if type(value) in (int, bool, str) and str(value) in (
                    "0", "1", "True", "False") else "非 0/1 状态，暂不解释"
        if not isinstance(layer.get("review"), dict):
            break
        layer, path = layer["review"], path + ".review"
    selected = like_state(unwrap_review_detail(data), official_detail=True)
    report = {"has_text": bool(text_content(row)), "has_chapter": row.get("chapterUid") is not None,
              "has_author_id": bool(author_id(row)), "state_fields": states, "selected_state": selected}
    print("\n单条想法验证（不显示正文和凭证）：")
    print("  正文可读取：" + ("是" if report["has_text"] else "否"))
    print("  含章节位置：" + ("是" if report["has_chapter"] else "否"))
    print("  含作者身份：" + ("是" if report["has_author_id"] else "否"))
    if states:
        for key, value in states.items():
            print(f"  {key} = {value}")
        print("  详情正文层级的状态：" + ("已点赞" if selected is True else "未点赞" if selected is False else "未知或冲突"))
        print("  官方单条详情的 isLike 已通过点赞为 1、取消点赞为 0 的对照验证。")
    else:
        print("  官方单条接口未返回 isLiked / isLike，不能据此判断您是否点赞。")
    return report


def like_state(row: dict, *, official_detail: bool = False) -> bool | None:
    """缺字段、异常值或两个状态互相矛盾都视为未知，不根据赞数推断。

    isLike 仅在已实测的官方 /review/single 详情上解释；旧列表不自动套用。
    """
    values = []
    for key in (("isLiked", "isLike") if official_detail else ("isLiked",)):
        if key in row:
            value = row[key]
            if value is True or type(value) is int and value == 1 or value == "1":
                values.append(True)
            elif value is False or type(value) is int and value == 0 or value == "0":
                values.append(False)
            else:
                return None
    return values[0] if values and len(set(values)) == 1 else None


def number(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except (ValueError, TypeError, OverflowError):
        return default


def clean(value: Any) -> str:
    """去除不能写进 OOXML 的控制字符，保留中文、换行和 Tab。"""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", str(value or ""))


def text_content(row: dict) -> str:
    if row.get("content"):
        return clean(row["content"])
    raw = re.sub(r"<br\s*/?>|</p>", "\n", str(row.get("htmlContent", "")), flags=re.I)
    return clean(html.unescape(re.sub(r"<[^>]*>", "", raw))).strip()


def date_text(value: Any, timezone: ZoneInfo) -> str:
    try:
        timestamp = float(value)
        if timestamp <= 0:
            return "日期未知"
        if timestamp > 100_000_000_000:  # 兼容毫秒时间戳
            timestamp /= 1000
        return datetime.fromtimestamp(timestamp, timezone).strftime("%Y-%m-%d %H:%M:%S %z")
    except (ValueError, TypeError, OverflowError, OSError):
        return "日期未知"


def unwrap_review(item: dict) -> dict:
    """兼容 review、review.review 包装，保留外层 isLiked / idx 等元数据。

    同名字段优先取最外层，避免把内层过期状态覆盖列表中的当前点赞状态。
    """
    layers, row = [], item
    for _ in range(6):
        layers.append(row)
        if not isinstance(row.get("review"), dict):
            break
        row = row["review"]
    result: dict = {}
    for layer in reversed(layers):
        result.update({k: v for k, v in layer.items() if k != "review"})
    return result


def unwrap_review_detail(item: dict) -> dict:
    """单条详情以最内层想法为准，与列表包装对象的元数据优先级分开处理。

    外层的默认 isLiked=0 不应与实际 review.isLike=1 混为一条互相矛盾的状态。
    同一层内两个状态冲突仍由 like_state 判为未知，绝不强制视作已点赞。
    """
    result, flags = {}, {}
    layer = item
    for _ in range(6):
        current_flags = {k: layer[k] for k in ("isLiked", "isLike") if k in layer}
        if current_flags:
            flags = current_flags
        result.update({k: v for k, v in layer.items()
                       if k not in ("review", "isLiked", "isLike") and v is not None})
        if not isinstance(layer.get("review"), dict):
            break
        layer = layer["review"]
    result.update(flags)
    return result


def merge_review_context(context: dict, row: dict) -> dict:
    """空的可选字段不覆盖已知的章节、原文范围和同一 ID 的列表元数据。"""
    result = dict(context)
    for key, value in row.items():
        if value is not None and value != "" and value != {}:
            result[key] = value
    return result


def author_id(row: dict) -> str:
    author = row.get("author") if isinstance(row.get("author"), dict) else {}
    return str(author.get("userVid") or author.get("vid") or row.get("userVid") or "")


def unique_rows(rows: list[dict], id_key: str) -> list[dict]:
    # 相同内容、不同 ID 的记录不合并；无 ID 时才退回整个记录的摘要。
    result, seen = [], set()
    for row in rows:
        key = str(row.get(id_key) or hashlib.sha256(
            json.dumps(row, sort_keys=True, ensure_ascii=False).encode()).hexdigest())
        if key not in seen:
            seen.add(key)
            result.append(row)
    return result


def response_shape(value: Any, depth: int = 0) -> str:
    """只描述字段名、类型、数组长度，绝不显示字符串内容、数值或列表记录。

    出错时可直接截图此诊断；不落盘原始 API 响应或个人笔记。
    """
    if isinstance(value, dict):
        if depth >= 2:
            return "object"
        fields = []
        for key, item in list(value.items())[:20]:
            name = str(key)
            if not re.fullmatch(r"[A-Za-z_][A-Za-z_0-9]{0,39}", name):
                name = "[字段名已隐藏]"
            fields.append(f"{name}:{response_shape(item, depth + 1)}")
        return "{" + ", ".join(fields) + (", …" if len(value) > 20 else "") + "}"
    if isinstance(value, list):
        return f"array({len(value)})"
    return {str: "string", int: "number", float: "number", bool: "boolean",
            type(None): "null"}.get(type(value), "other")


def extract_highlights(data: dict) -> tuple[list[dict], dict] | None:
    """只解析明确的 updated 数组；不因字段缺失就猜测没有笔记。"""
    current = data
    for _ in range(4):
        if "updated" in current:
            rows = current["updated"]
            if isinstance(rows, list) and all(isinstance(row, dict) for row in rows):
                return rows, current
            return None
        if isinstance(current.get("data"), dict):
            current = current["data"]
        else:
            break
    # 空集必须有明确的数量信息；{}、null、只有 synckey 都不能当成空集。
    if str(current.get("totalCount")) == "0":
        return [], current
    return None


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        # 登录重定向应提示 Cookie 失效，也避免凭证被转发到其它主机。
        return None


class Weread:
    WEB = "https://weread.qq.com"
    GATEWAY = "https://i.weread.qq.com/api/agent/gateway"
    LEGACY = "https://i.weread.qq.com"

    def __init__(self, cookie: str = "", api_key: str = "", max_pages: int = 1000,
                 interactive_auth: bool = False, personal_only: bool = False, show_dates: bool = False,
                 liked_source: str = "paragraphs", review_ids: list[str] | None = None,
                 max_review_checks: int = 200, max_review_ranges: int = 500):
        self.cookie, self.vid = parse_cookie(cookie) if cookie else ("", "")
        self.api_key, self.max_pages = api_key, max_pages
        self.interactive_auth = interactive_auth
        self.personal_only = personal_only
        self.show_dates = show_dates
        self.liked_source, self.review_ids = liked_source, review_ids or []
        self.max_review_checks, self.max_review_ranges = max_review_checks, max_review_ranges
        self.cookie_policy = DefaultCookiePolicy()
        self.cookiejar = CookieJar(policy=self.cookie_policy)
        self._seeded_cookies = {}
        parsed = SimpleCookie()
        parsed.load(self.cookie)
        for name, morsel in parsed.items():
            item = Cookie(0, name, morsel.value, None, False, "weread.qq.com", False,
                          False, "/", True, True, None, True, None, None, {})
            self.cookiejar.set_cookie(item)
            self._seeded_cookies[name] = item
        self.opener = build_opener(HTTPCookieProcessor(self.cookiejar), NoRedirect())
        # Agent 请求不能共用携带网页 Cookie 的 opener。
        self.agent_opener = build_opener(NoRedirect())
        self.last_request = 0.0

    def _accept_rotated_cookies(self, response, request: Request):
        """CookieJar 处理 Set-Cookie；清除同名的手工引入副本，防止发送两份 skey。

        从 Cookie 请求头导入时不知道原始 Domain。服务器重新下发后，以服务端
        Domain/Path 为准。只有 CookieJar 实际接受的 Cookie 才能替换导入副本。
        """
        incoming = self.cookiejar.make_cookies(response, request)
        for item in incoming:
            seed = self._seeded_cookies.get(item.name)
            if (seed is not None and item.domain.lstrip(".") == seed.domain.lstrip(".")
                    and item.path == seed.path and self.cookie_policy.set_ok(item, request)):
                if (seed.domain, seed.path) != (item.domain, item.path):
                    try:
                        self.cookiejar.clear(seed.domain, seed.path, seed.name)
                    except KeyError:
                        pass
                self._seeded_cookies.pop(item.name, None)
        # 同样适用于 HTTPError；更新、删除和过期均由标准库按 Cookie 规则处理。
        self.cookiejar.extract_cookies(response, request)

    def refresh_session(self):
        """只做一次同源 HEAD 握手，吸收服务端下发的新 Cookie；不代表登录已通过。"""
        request = Request(self.WEB + "/", method="HEAD", headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
            "Accept": "*/*", "Referer": self.WEB + "/"})
        try:
            with self.opener.open(request, timeout=30) as response:
                self._accept_rotated_cookies(response, request)
        except HTTPError as exc:
            self._accept_rotated_cookies(exc, request)
            if exc.code not in (301, 302, 303, 307, 308, 405):
                raise ExportError(f"会话更新握手失败（HTTP {exc.code}），请检查网络或登录状态。") from None
            # 不跟随重定向或把 Cookie 发送到其它站点；仍尝试原接口验证结果。
        except (URLError, TimeoutError, ConnectionError):
            raise ExportError("会话更新时网络连接失败，请检查网络后重试。") from None

    def _session_failure(self, path: str, status: str = "-2012") -> ExportError:
        if path != "/api/user/notebook":
            try:
                probe = self.request("/api/user/notebook", _recover_session=False)
            except ExportError:
                pass
            else:
                if isinstance(probe.get("books"), list):
                    return ExportError(f"{path} 在会话更新后仍返回 {status}，但个人笔记接口验证成功。"
                                       "当前凭证可访问个人笔记，此接口可能有额外认证要求或已变化；"
                                       "请保留此错误以继续排查，无需反复复制 Cookie。")
        return ExportError(f"{path} 在会话更新并重试后仍返回 {status}，个人登录状态未能验证。"
                           "可能是凭证失效、缺少会话字段或服务端拒绝此客户端；"
                           "仅凭此代码无法确定原因，请保留错误以继续排查。")

    def request(self, path: str, params: dict | None = None, *, post: bool = False,
                agent: bool = False, legacy: bool = False, _recover_session: bool = True) -> dict:
        # 旧接口只开放已核对的划线查询，不能传入任意主机或任意路径。
        if legacy and (agent or path != "/book/bookmarklist"):
            raise ExportError("不支持的兼容接口请求。")
        params = params or {}
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36", "Accept": "application/json",
                   "Content-Type": "application/json", "Accept-Encoding": "identity"}
        if agent:
            url, post = self.GATEWAY, True
            payload = {"api_name": path, "skill_version": "1.0.5", **params}
            headers["Authorization"] = "Bearer " + self.api_key
        else:
            if not self.cookie:
                raise ExportError("读取已点赞他人想法需要同一账号的 Cookie；Agent API 未公开等价能力。")
            url, payload = (self.LEGACY if legacy else self.WEB) + path, params
            headers["Referer"] = self.WEB + "/"
        if not post and params:
            url += "?" + urlencode(params)
        body = json.dumps(payload).encode() if post else None
        session_retried = False
        for attempt in range(5):
            time.sleep(max(0, 0.4 - (time.monotonic() - self.last_request)))
            self.last_request = time.monotonic()
            # 每次构造新 Request，让 CookieJar 注入最新凭证，绝不重发固定 Cookie 头。
            request = Request(url, data=body, headers=headers)
            try:
                opener = self.agent_opener if agent else self.opener
                with opener.open(request, timeout=30) as response:
                    if not agent:
                        self._accept_rotated_cookies(response, request)
                    data = json.loads(response.read().decode("utf-8"))
            except HTTPError as exc:
                if not agent:
                    self._accept_rotated_cookies(exc, request)
                if exc.code == 401 and not agent and _recover_session:
                    if not session_retried and attempt < 4:
                        print("  登录请求被拒绝，正在更新会话并重试一次…")
                        self.refresh_session()
                        session_retried = True
                        continue
                    raise self._session_failure(path, "HTTP 401") from None
                if exc.code in (429, 500, 502, 503, 504) and attempt < 3:
                    delay = max(2 ** attempt, number(exc.headers.get("Retry-After"), 0))
                    print(f"  服务繁忙，{delay} 秒后重试…")
                    time.sleep(delay)
                    continue
                if exc.code in (301, 302, 303, 307, 308, 401, 403):
                    raise ExportError(f"{path} 登录或授权失败，请更新 Cookie / API Key。") from None
                raise ExportError(f"{path} 请求失败（HTTP {exc.code}）。") from None
            except (URLError, TimeoutError, ConnectionError):
                if attempt < 3:  # 微信请求均为只读，含查询型 POST，可以安全重试
                    time.sleep(2 ** attempt)
                    continue
                raise ExportError(f"{path} 网络连接失败，请检查网络后重试。") from None
            except (ValueError, UnicodeError):
                raise ExportError(f"{path} 未返回有效 JSON，可能需要重新登录或接口已变化。") from None
            if not isinstance(data, dict):
                raise ExportError(f"{path} 返回结构已变化：预期 JSON 对象。")
            if data.get("upgrade_info"):
                raise ExportError("Agent Gateway 要求升级，请核对官方 Skill 版本并更新 skill_version。")
            # 先展开通用 data 对象，再检查业务错误，避免把内层错误当成缺字段。
            for _ in range(4):
                outer_code = data.get("errcode", data.get("errCode", 0))
                if str(outer_code) not in ("0", "None", "") or not isinstance(data.get("data"), dict):
                    break
                data = data["data"]
                if data.get("upgrade_info"):
                    raise ExportError("Agent Gateway 要求升级，请核对官方 Skill 版本并更新 skill_version。")
            code = data.get("errcode", data.get("errCode", 0))
            if str(code) == "-2012" and not agent and _recover_session:
                if not session_retried and attempt < 4:
                    print("  收到 -2012，正在接收更新 Cookie、刷新会话并重试一次…")
                    self.refresh_session()
                    session_retried = True
                    continue
                raise self._session_failure(path)
            if str(code) not in ("0", "None", ""):
                raise ExportError(f"{path} 业务错误（代码 {number(code)}），请检查登录状态和接口支持情况。")
            return data
        raise ExportError("请求重试次数耗尽。")

    def pages(self, path: str, params: dict, key: str, *, agent: bool = False) -> list[dict]:
        """游标翻页。hasMore 未清零、游标停滞、重复页时明确失败，不静默截断。"""
        result, seen_pages = [], set()
        query = dict(params)
        for index in range(self.max_pages):
            data = self.request(path, query, agent=agent)
            rows = data.get(key)
            if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
                raise ExportError(f"{path} 缺少 {key} 数组，请检查接口结构。")
            digest = hashlib.sha256(json.dumps(rows, sort_keys=True).encode()).hexdigest()
            more_value = data.get("hasMore", data.get("reviewsHasMore"))
            more = str(more_value).lower() in ("1", "true")
            if rows and digest in seen_pages:
                raise ExportError(f"{path} 返回重复页，无法确认是否完整，请勿将当前结果当作全量。")
            seen_pages.add(digest)
            result.extend(rows)
            print(f"  已读取第 {index + 1} 页，累计 {len(result)} 条")
            if not more:
                total = data.get("totalCount")
                if more_value is None and total is not None and number(total) > len(result):
                    raise ExportError(f"{path} 存在未读取记录，但没有可识别的翻页标记。")
                return result
            if not rows:
                raise ExportError(f"{path} 声明仍有下一页，却返回空页。")
            updated = dict(query)
            if "synckey" in data or "syncKey" in data:
                updated["synckey"] = data.get("synckey", data.get("syncKey"))
            if "maxIdx" in data:
                updated["maxIdx"] = data["maxIdx"]
            elif "idx" in rows[-1]:
                updated["maxIdx"] = rows[-1]["idx"]
            if updated == query:
                raise ExportError(f"{path} 分页游标没有前进，停止以免遗漏或死循环。")
            query = updated
        raise ExportError(f"{path} 达到 {self.max_pages} 页保护上限，尚未读取完毕。")

    def shelf(self) -> list[dict]:
        # 真正的书架，不用“有笔记的书籍列表”替代书架。
        data = self.request("/shelf/sync" if self.api_key else "/web/shelf/sync",
                            {} if self.api_key else {"userVid": self.vid, "synckey": 0},
                            agent=bool(self.api_key))
        if str(data.get("hasMore", 0)).lower() in ("1", "true"):
            raise ExportError("书架接口返回分段同步结果，当前适配器不能确认全量书架，请更新适配器。")
        rows = data.get("books")
        if not isinstance(rows, list):
            raise ExportError("书架响应缺少 books，请检查凭证或接口变化。")
        books = []
        for item in rows:
            if not isinstance(item, dict):
                continue
            row = item.get("book") if isinstance(item.get("book"), dict) else item
            if row.get("bookId") and not row.get("deleted"):
                books.append(row)
        extra = len(data.get("albums") or []) + bool(data.get("mp"))
        if extra:
            print(f"  书架另有 {extra} 个有声书/文章集合，本工具仅列出支持笔记导出的 books 条目。")
        return unique_rows(books, "bookId")

    def collect(self, book: dict, timezone: ZoneInfo, min_likes: int) -> "Notebook":
        bid = str(book["bookId"])
        agent = bool(self.api_key)
        print("[3/5] 正在抓取书籍信息与章节目录…")
        info = self.request("/book/info" if agent else "/web/book/info", {"bookId": bid}, agent=agent)
        info = info.get("book", info)
        if not isinstance(info, dict) or not info.get("title"):
            raise ExportError("书籍详情缺少书名，无法可靠导出。")
        if agent:
            catalog = self.request("/book/chapterinfo", {"bookId": bid}, agent=True)
        else:
            catalog = self.request("/web/book/chapterInfos", {"bookIds": [bid]}, post=True)
        print("[3/5] 正在抓取个人划线…")
        highlights, marks = self.highlights(bid, agent=agent)
        agent = bool(self.api_key)  # 划线读取阶段可能已由用户切换到 Gateway。
        if str(marks.get("hasMore", 0)).lower() in ("1", "true"):
            raise ExportError("划线接口出现分页；当前适配器无法确认全量，请更新适配器。")
        chapters = extract_chapters(catalog, bid)
        # bookmarklist 的附带章节是目录的补充来源。
        known = {str(c.get("chapterUid")) for c in chapters}
        chapters += [c for c in marks.get("chapters", []) if str(c.get("chapterUid")) not in known]
        print("[3/5] 正在抓取个人想法…")
        mine = self.pages("/review/list/mine" if agent else "/web/review/list",
                          {"bookid": bid, "count": 100, "synckey": 0} if agent else
                          {"bookId": bid, "listType": 11, "mine": 1, "count": 100, "synckey": 0},
                          "reviews", agent=agent)
        mine = [unwrap_review(r) for r in mine]
        own_authors = {author_id(row) for row in mine if author_id(row)}
        if len(own_authors) > 1 or self.vid and any(v != self.vid for v in own_authors):
            raise ExportError("个人想法中的账号身份不一致，请使用同一账号的凭证。")
        own_vid = self.vid or next(iter(own_authors), "")
        own_ids = {str(row["reviewId"]) for row in mine if row.get("reviewId")}
        mine_ids_complete = all(row.get("reviewId") for row in mine)
        # Gateway 的公开点评接口不保证提供 isLiked，不能用它冒充“我的点赞”。
        warnings = []
        public = []
        public_available = False
        if self.personal_only:
            warnings.append("仅个人笔记导出：本次按用户选择未抓取已点赞他人评论，文档仅含个人划线与想法。")
        else:
            print("[3/5] 正在读取评论并严格筛选我点赞的他人想法…")
            try:
                if self.api_key:
                    public, scope_warnings = self.liked_thoughts(bid, chapters, highlights + mine, own_ids)
                    warnings.extend(scope_warnings)
                else:
                    public = self.pages("/web/review/list", {"bookId": bid, "listType": 11, "mine": 0,
                                        "count": 100, "maxIdx": 0, "synckey": 0}, "reviews")
                public_available = True
            except ExportError as exc:
                if not self.interactive_auth:
                    raise
                print(f"  已点赞评论读取失败：{exc}")
                print("  个人划线和想法已抓取，可以导出；但无法核验已点赞他人评论。")
                action = choice("请选择：1. 继续导出个人笔记（文档注明缺少点赞评论）；2. 停止导出 [默认 2]：",
                                ("1", "2"), "2")
                if action != "1":
                    raise ExportError("已停止导出：用户未选择缺少点赞评论的个人笔记导出。") from None
                warnings.append("已点赞他人评论接口读取失败：用户选择继续导出个人划线与想法，本文件不包含已点赞他人评论。")
        public = [unwrap_review(r) for r in public]
        if public_available and not self.api_key:
            warnings.append("点赞评论来自接口当前可见的 listType=11 评论范围；服务端不可见或未返回的评论无法核验。")
        missing_status = sum("isLiked" not in r for r in public)
        missing_author = sum(str(r.get("isLiked")) == "1" and not author_id(r) for r in public)
        if missing_status:
            warnings.append(f"{missing_status} 条评论未返回 isLiked，已排除，点赞评论结果可能不完整。")
        if missing_author:
            warnings.append(f"{missing_author} 条已点赞评论缺少作者身份，无法确认是他人，已排除。")
        if public_available and not public:
            warnings.append("评论接口返回空列表；这不能证明全书没有您点赞过的评论。")
        liked, excluded = [], defaultdict(int)
        explicit_ids = {review_id_from_reference(reference) for reference in self.review_ids}
        for row in public:
            rid = str(row.get("reviewId") or "")
            reason = ""
            if like_state(row) is not True:
                reason = "点赞状态未确认"
            elif not author_id(row):
                reason = "缺少作者身份"
            elif rid in own_ids or own_vid and author_id(row) == own_vid:
                reason = "属于本人想法"
            elif not own_vid and not (mine_ids_complete and rid):
                reason = "无法排除本人内容"
            elif number(row.get("likesCount", row.get("likeCount"))) < min_likes:
                reason = "低于赞数门槛或未返回赞数"
            elif not text_content(row):
                reason = "正文为空"
            if reason:
                excluded[reason] += 1
                if rid in explicit_ids:
                    print(f"  指定想法 {rid} 未收录：{reason}。")
            else:
                liked.append(row)
                if rid in explicit_ids:
                    print(f"  指定想法 {rid} 已通过最终筛选，将写入读者观点。")
        if excluded:
            print("  最终筛选排除原因：" + "；".join(f"{key} {value} 条" for key, value in excluded.items()))
            warnings.append("部分候选观点未收录：" + "；".join(f"{key} {value} 条" for key, value in excluded.items()) + "。")
        if public and not own_vid and not mine_ids_complete:
            warnings.append("当前账号身份无法确认且个人想法缺少 ID，为避免混入本人评论，未收录读者观点。")
        return build_notebook({**book, **info, "bookId": bid}, chapters,
                              highlights, mine, liked, timezone, warnings, show_dates=self.show_dates)

    def liked_thoughts(self, bid: str, chapters: list[dict], highlights: list[dict],
                      own_ids: set[str]) -> tuple[list[dict], list[str]]:
        """发现候选后逐条用官方 single 核验，范围有限，始终写明覆盖边界。

        分享链接优先，不依赖热门排名；自动发现来自自己的划线以及各章 TOP20。
        TOP20 之外、没有已知 range 的段落不会被假装已扫描。
        """
        candidates: dict[str, dict] = {}
        warnings = []
        explicit_ids = {review_id_from_reference(reference) for reference in self.review_ids}
        priorities = {}
        for reference in self.review_ids:
            rid = review_id_from_reference(reference)
            candidates.setdefault(rid, {})
            priorities[rid] = 0
        if len(candidates) > self.max_review_checks:
            raise ExportError("提供的想法链接数超过 --max-review-checks，请提高上限或分批验证。")
        ranges: dict[tuple[str, str], dict] = {}
        discovery_limit = self.max_review_checks * 5
        failed = 0
        def add_range(row):
            uid, span = str(row.get("chapterUid", "")), str(row.get("range") or "")
            if uid and re.fullmatch(r"\d+-\d+", span):
                ranges.setdefault((uid, span), {"abstract": clean(row.get("markText") or row.get("abstract")), "chapterUid": uid, "range": span})
        seed_details = {}
        if self.liked_source == "paragraphs":
            warnings.append("读者观点仅覆盖个人笔记、指定链接及各章热门划线中可定位的段落，不代表全书完整点赞记录。")
            # 用户给出的链接也是原文定位入口，可补充热门 TOP20 之外的段落。
            for rid in candidates:
                if rid in own_ids:
                    continue
                data = self.request("/review/single", {"reviewId": rid}, agent=True)
                row = unwrap_review_detail(data)
                if str(row.get("reviewId") or "") != rid:
                    raise ExportError("指定想法详情的 ID 与请求不一致。")
                book = row.get("book") if isinstance(row.get("book"), dict) else {}
                book_id = str(row.get("bookId") or book.get("bookId") or "")
                if book_id and book_id != bid:
                    raise ExportError("提供的想法链接属于另一本书，请在书架中选择对应书籍。")
                seed_details[rid] = data
                if book_id == bid:
                    add_range(row)
            for mark in highlights:
                add_range(mark)
            seen_chapters = set()
            for chapter in chapters:
                uid = str(chapter.get("chapterUid", ""))
                if not uid or uid in seen_chapters:
                    continue
                seen_chapters.add(uid)
                if len(ranges) >= self.max_review_ranges:
                    warnings.append("段落扫描达到本次上限，剩余章节未扫描。")
                    break
                print(f"  正在查找第 {len(seen_chapters)} 个章节中的读者想法…")
                try:
                    data = self.request("/book/bestbookmarks", {"bookId": bid, "chapterUid": number(uid)}, agent=True)
                    if not isinstance(data.get("items"), list):
                        raise ExportError("热门划线响应缺少 items。")
                    for item in data["items"]:
                        if isinstance(item, dict):
                            add_range({**item, "chapterUid": item.get("chapterUid", uid)})
                except ExportError:
                    failed += 1
            page_count = 0
            for index, ((uid, span), context) in enumerate(ranges.items()):
                if index >= self.max_review_ranges or len(candidates) >= discovery_limit:
                    warnings.append("候选想法或段落扫描达到本次上限，仍可能有未发现的点赞内容。")
                    break
                if index % 10 == 0:
                    print(f"  正在读取原文范围 {index + 1}/{min(len(ranges), self.max_review_ranges)}，已发现 {len(candidates)} 条候选想法…")
                cursor, seen = {"maxIdx": 0, "synckey": 0}, set()
                for _ in range(self.max_pages):
                    if page_count >= self.max_pages:
                        warnings.append("读者想法扫描达到请求页数上限，本次结果不完整。")
                        break
                    page_count += 1
                    try:
                        data = self.request("/book/readreviews", {"bookId": bid, "chapterUid": number(uid),
                            "reviews": [{"range": span, "count": 20, **cursor}]}, agent=True)
                        groups = data.get("reviews")
                        if not isinstance(groups, list):
                            raise ExportError("段落想法响应缺少 reviews。")
                        matches = [r for r in groups if isinstance(r, dict) and str(r.get("range")) == span]
                        if len(matches) != 1 or not isinstance(matches[0].get("pageReviews"), list):
                            raise ExportError("段落想法响应未包含对应 range。")
                        group = matches[0]
                        rows = group["pageReviews"]
                        if "hasMore" not in group and number(group.get("totalCount"), -1) > len(rows):
                            raise ExportError("段落想法有更多记录，但缺少分页标记。")
                        signature = hashlib.sha256(json.dumps(rows, sort_keys=True).encode()).hexdigest()
                        if signature in seen:
                            raise ExportError("段落想法分页重复。")
                        seen.add(signature)
                        for entry in rows:
                            if not isinstance(entry, dict):
                                continue
                            row = unwrap_review(entry)
                            rid = str(row.get("reviewId") or "")
                            if rid and rid not in own_ids:
                                metadata = {k: row[k] for k in ("chapterUid", "chapterName", "range", "abstract", "author", "userVid", "likesCount", "likeCount") if k in row}
                                detail_context = merge_review_context(context, metadata)
                                candidates[rid] = merge_review_context(detail_context, candidates.get(rid, {}))
                                # 列表状态只影响核验顺序，绝不用于认定已点赞；每条仍查询 single。
                                hint = like_state(row, official_detail=True)
                                priority = 1 if hint is True else 3 if hint is False else 2
                                priorities[rid] = min(priorities.get(rid, priority), priority)
                                if len(candidates) >= discovery_limit:
                                    warnings.append("候选想法达到本次上限，部分内容可能未扫描。")
                                    break
                        if str(group.get("hasMore", 0)).lower() not in ("1", "true"):
                            break
                        next_cursor = {k: group.get(k, cursor[k]) for k in cursor}
                        if next_cursor == cursor or not rows:
                            raise ExportError("段落想法分页游标停滞。")
                        if len(candidates) >= discovery_limit:
                            warnings.append("候选想法达到本次上限，部分内容未扫描。")
                            break
                        cursor = next_cursor
                    except ExportError:
                        failed += 1
                        break
        else:
            if not candidates:
                raise ExportError("链接模式需要至少一个 --review-id 想法链接或 ID。")
            warnings.append("读者观点仅核验本次提供的想法链接，未扫描全书其他点赞记录。")
        result, unknown, unliked, verified = [], 0, 0, 0
        ordered = sorted(candidates.items(), key=lambda item: priorities.get(item[0], 2))
        if len(ordered) > self.max_review_checks:
            warnings.append(f"发现 {len(ordered)} 条候选想法，本次仅核验前 {self.max_review_checks} 条，仍可能遗漏；可提高 --max-review-checks。")
        print(f"  共发现 {len(candidates)} 条候选想法；指定链接优先，其次核验列表提示已点赞的内容。")
        for index, (rid, context) in enumerate(ordered[:self.max_review_checks]):
            if rid in own_ids:
                if rid in explicit_ids:
                    print(f"  指定想法 {rid} 未收录：个人笔记接口将其列为本人内容。")
                continue
            if index % 10 == 0:
                print(f"  正在逐条核验点赞状态 {index + 1}/{min(len(candidates), self.max_review_checks)}…")
            try:
                data = seed_details[rid] if rid in seed_details else self.request("/review/single", {"reviewId": rid}, agent=True)
                row = unwrap_review_detail(data)
                if str(row.get("reviewId") or "") != rid:
                    raise ExportError("想法详情 ID 不一致。")
                book = row.get("book") if isinstance(row.get("book"), dict) else {}
                book_id = str(row.get("bookId") or book.get("bookId") or "")
                if book_id and book_id != bid:
                    if rid in explicit_ids:
                        raise ExportError("提供的想法链接属于另一本书，请在书架中选择对应书籍。")
                    continue
                if not book_id and not context:
                    raise ExportError("想法缺少书籍信息，不能关联到当前书籍。")
                state = like_state(row, official_detail=True)
                verified += 1
                if state is None:
                    unknown += 1
                    if rid in explicit_ids:
                        print(f"  指定想法 {rid} 未收录：点赞字段缺失、异常或同一层级状态冲突。")
                elif state is False:
                    unliked += 1
                    if rid in explicit_ids:
                        print(f"  指定想法 {rid} 未收录：官方详情当前返回未点赞（状态 0）。")
                elif text_content(row):
                    result.append({**merge_review_context(context, row), "isLiked": 1})
                    if rid in explicit_ids:
                        print(f"  指定想法 {rid} 已确认点赞，正在核对作者与导出条件。")
                elif rid in explicit_ids:
                    print(f"  指定想法 {rid} 未收录：详情正文为空。")
            except ExportError:
                # 用户指定的样本失败时必须明确停止，不能静默跳过验证目标。
                if rid in explicit_ids:
                    raise
                failed += 1
        if failed:
            warnings.append(f"{failed} 个章节、段落或想法未能读取，本次读者观点不完整。")
        if unknown:
            warnings.append(f"{unknown} 条想法没有可确认的点赞状态，未收录。")
        if failed and not verified:
            raise ExportError(f"读者想法扫描有 {failed} 处读取失败，未能核验任何候选想法，不能将此结果视为没有点赞记录。")
        print(f"  已核验 {verified} 条想法：已点赞 {len(result)} 条，未点赞 {unliked} 条，状态未知 {unknown} 条。")
        return unique_rows(result, "reviewId"), list(dict.fromkeys(warnings))

    def highlights(self, bid: str, *, agent: bool = False) -> tuple[list[dict], dict]:
        # 网页接口是增量同步接口，显式从 0 请求，避免默认值造成无更新回包。
        params: dict = {"bookId": bid}
        reported_count = None
        if not agent:
            params["synckey"] = 0
        data = self.request("/book/bookmarklist" if agent else "/web/book/bookmarklist", params, agent=agent)
        parsed = extract_highlights(data)
        if parsed is not None:
            return parsed
        print("  划线接口返回结构：" + response_shape(data))
        # 有些空笔记回包省略 updated。用个人笔记概览的明确 noteCount 二次核验，
        # 不是用概览替代书架，也不会因书籍不在概览里就推断划线数为零。
        if not agent:
            try:
                overview = self.request("/api/user/notebook")
            except ExportError:
                print("  未能通过个人笔记概览核验划线数量。")
            else:
                books = overview.get("books")
                if isinstance(books, list):
                    for entry in books:
                        if not isinstance(entry, dict):
                            continue
                        info = entry.get("book") if isinstance(entry.get("book"), dict) else {}
                        if str(entry.get("bookId") or info.get("bookId")) != bid:
                            continue
                        if str(entry.get("noteCount")) == "0":
                            print("  个人笔记概览确认该书划线数为 0，将继续抓取个人想法和点赞评论。")
                            return [], data
                        if number(entry.get("noteCount"), -1) > 0:
                            reported_count = number(entry["noteCount"])
                        break
            # 部分网页版本需要 type=1；另有公开项目使用 i.weread.qq.com 的同名接口。
            # 两条回退路径都只查本人的 bookmarklist，不用热门划线冒充个人划线。
            variants = [
                ("网页 type=1", "/web/book/bookmarklist", {"bookId": bid, "synckey": 0, "type": 1}, False),
                ("i.weread.qq.com 兼容接口", "/book/bookmarklist", {"bookId": bid, "synckey": 0}, True),
            ]
            for label, path, query, legacy in variants:
                print(f"  正在尝试{label}读取个人划线…")
                try:
                    alternative = self.request(path, query, legacy=legacy)
                except ExportError as exc:
                    print(f"  {label}未成功：{exc}")
                    continue
                parsed = extract_highlights(alternative)
                if parsed is not None:
                    actual = sum(str(row.get("type", 1)) == "1" and bool(row.get("markText")) for row in parsed[0])
                    if reported_count and actual < reported_count:
                        print(f"  {label}划线条数少于概览，不能确认全量，继续尝试。")
                        continue
                    print(f"  已通过{label}读取到 {actual} 条个人划线。")
                    return parsed
                print(f"  {label}返回结构：{response_shape(alternative)}")
            if self.interactive_auth:
                print("  Cookie 划线接口不可用，可以改用微信读书官方 Agent API。")
                print("  请用相同账号打开 https://weread.qq.com/r/weread-skills 获取 API Key。")
                action = choice("请选择：1. 输入 Agent API Key 继续当前书籍；2. 停止导出 [默认 2]：", ("1", "2"), "2")
                if action == "1":
                    key = secret("请输入 Agent API Key（隐藏输入）：", "WEREAD_API_KEY", validator=validate_api_key)
                    previous = self.api_key
                    self.api_key = key
                    try:
                        parsed = self.highlights(bid, agent=True)
                        actual = sum(str(row.get("type", 1)) == "1" and bool(row.get("markText")) for row in parsed[0])
                        if reported_count and actual < reported_count:
                            raise ExportError("官方接口划线数少于当前 Cookie 账号概览，请确认 API Key 与 Cookie 属于同一账号。")
                    except Exception:
                        self.api_key = previous
                        raise
                    print(f"  已通过官方 Agent API 读取到 {actual} 条个人划线，后续个人想法也使用官方接口。")
                    return parsed
            if reported_count:
                raise ExportError("个人笔记概览显示该书有划线，但网页及兼容划线接口都未返回可确认的完整划线。"
                                  "已停止导出以免遗漏；请提供以上结构诊断和错误信息。")
        raise ExportError("划线接口未返回可识别的 updated 数组，暂时无法确认是否为空笔记。"
                          "已停止导出以免遗漏；请提供上方“划线接口返回结构”诊断，无需提供 Cookie。")


def extract_chapters(data: dict, bid: str) -> list[dict]:
    if isinstance(data.get("chapters"), list):
        return data["chapters"]
    for item in data.get("data", []) if isinstance(data.get("data"), list) else []:
        if str(item.get("bookId")) == bid:
            for key in ("updated", "chapterInfos", "chapters"):
                if isinstance(item.get(key), list):
                    return item[key]
    if isinstance(data.get("chapterInfos"), list):
        return data["chapterInfos"]
    raise ExportError("章节响应缺少 chapters / data[].updated / chapterInfos，不能确认目录结构。")


@dataclass
class Note:
    kind: str
    uid: str
    text: str
    date: str
    range: str = ""
    source_id: str = ""
    quote: str = ""


@dataclass
class Group:
    highlight: Note | None = None
    thoughts: list[Note] = field(default_factory=list)
    quote: str = ""  # 没有个人划线时，仍可把同一原文下的想法和观点归为一组。


@dataclass
class Chapter:
    uid: str
    title: str
    order: int
    groups: list[Group] = field(default_factory=list)


@dataclass
class Notebook:
    book_id: str
    title: str
    author: str
    chapters: list[Chapter]
    warnings: list[str] = field(default_factory=list)
    show_dates: bool = False


def position(value: str) -> int:
    match = re.match(r"^(\d+)", value)
    return int(match[1]) if match else 10 ** 18


def build_notebook(info: dict, chapters: list[dict], marks: list[dict], mine: list[dict],
                   liked: list[dict], timezone: ZoneInfo, warnings: list[str] | None = None,
                   show_dates: bool = False) -> Notebook:
    """按章节实际顺序聚合。仅在章节和 ID/range/原文都能可靠关联时附到划线。"""
    chapter_map = {}
    for index, c in enumerate(chapters):
        uid = str(c.get("chapterUid", ""))
        if uid and uid not in chapter_map:
            chapter_map[uid] = Chapter(uid, clean(c.get("title") or c.get("chapterName") or f"章节 {uid}"),
                                       number(c.get("chapterIdx"), index))
    chapter_map["__unknown__"] = Chapter("__unknown__", "未分类章节", 10 ** 18)
    title_uids: dict[str, list[str]] = defaultdict(list)
    for c in chapter_map.values():
        title_uids[c.title].append(c.uid)

    def make_note(row: dict, kind: str) -> Note:
        uid = str(row.get("chapterUid", ""))
        if uid not in chapter_map:
            candidates = title_uids.get(clean(row.get("chapterName")), [])
            uid = candidates[0] if len(candidates) == 1 else "__unknown__"
        return Note(kind, uid, clean(row.get("markText")) if kind == "highlight" else text_content(row),
                    date_text(row.get("createTime"), timezone), str(row.get("range") or ""),
                    str(row.get("bookmarkId") if kind == "highlight" else row.get("reviewId") or ""),
                    clean(row.get("abstract")))

    for row in unique_rows(marks, "bookmarkId"):
        if str(row.get("type", 1)) != "1" or not row.get("markText"):
            continue  # type=0 是位置书签，不是划线
        note = make_note(row, "highlight")
        chapter_map[note.uid].groups.append(Group(note))
    for kind, rows in (("thought", mine), ("liked", liked)):
        for row in unique_rows(rows, "reviewId"):
            note = make_note(row, kind)
            if not note.text:
                continue
            chapter = chapter_map[note.uid]
            candidates = []
            # 未分类章节中的内容不强行配对，防止跨章节同文误关联。
            if note.uid != "__unknown__":
                highlights = [g for g in chapter.groups if g.highlight]
                bookmark_id = str(row.get("bookmarkId") or "")
                if bookmark_id:
                    candidates = [g for g in highlights if g.highlight.source_id == bookmark_id]
                if not candidates and note.range:
                    candidates = [g for g in highlights if g.highlight.range == note.range]
                if not candidates and note.quote:
                    candidates = [g for g in highlights if g.highlight.text.strip() == note.quote.strip()
                                  and not (note.range and g.highlight.range and note.range != g.highlight.range)]
                if not candidates:
                    # 同章、同一准确范围或唯一原文才能合并，不凭文字相似度猜关联。
                    for g in chapter.groups:
                        if g.highlight or not g.thoughts:
                            continue
                        source = g.thoughts[0]
                        same_range = bool(note.range and source.range == note.range)
                        same_quote = bool(note.quote and g.quote.strip() == note.quote.strip())
                        compatible = not (note.quote and g.quote and not same_quote)
                        if compatible and (same_range or same_quote and not (note.range and source.range and not same_range)):
                            candidates.append(g)
            if len(candidates) == 1:
                candidates[0].thoughts.append(note)
                if not candidates[0].highlight and not candidates[0].quote:
                    candidates[0].quote = note.quote
            else:
                chapter.groups.append(Group(thoughts=[note], quote=note.quote))
    result = []
    for chapter in sorted(chapter_map.values(), key=lambda c: c.order):
        if not chapter.groups:
            continue
        chapter.groups.sort(key=lambda g: (position((g.highlight or g.thoughts[0]).range),
                                           (g.highlight or g.thoughts[0]).date))
        for group in chapter.groups:
            group.thoughts.sort(key=lambda n: (n.kind == "liked", n.date))
        result.append(chapter)
    return Notebook(str(info["bookId"]), clean(info["title"]), clean(info.get("author") or "作者未知"),
                    result, warnings or [], show_dates)


def dated_text(note: Note, show_dates: bool) -> str:
    """仅在用户需要且日期有效时，在原文末尾显示短日期；不改动原文内容。"""
    date = re.match(r"^\d{4}-\d{2}-\d{2}", note.date)
    return note.text + (f"（{date[0]}）" if show_dates and date else "")


def display_notice(warning: str) -> str:
    """正文只保留与阅读有关的范围说明，接口排错详情留在终端。"""
    if "本次按用户选择未抓取" in warning:
        return "本次仅含个人划线与想法，未包含已点赞的读者观点。"
    if "已点赞他人评论接口读取失败" in warning:
        return "读者观点未能读取。本次保留个人划线与想法。"
    if "点赞评论来自接口当前可见" in warning:
        return "读者观点仅包含当前可读取、已确认由我点赞的内容。"
    if "条评论未返回 isLiked" in warning:
        count = re.match(r"\d+", warning)
        return f"{count[0] if count else '部分'}条读者观点无法确认点赞状态，未收录。"
    if "条已点赞评论缺少作者身份" in warning:
        count = re.match(r"\d+", warning)
        return f"{count[0] if count else '部分'}条评论无法确认作者身份，未收录。"
    if "评论接口返回空列表" in warning:
        return "本次未读取到读者观点，无法据此判断是否有过点赞记录。"
    return warning


def concise_notices(warnings: list[str]) -> list[str]:
    """阅读文档只留一条范围说明，完整错误和统计继续显示在终端。"""
    if not warnings:
        return []
    if any("本次按用户选择未抓取" in w for w in warnings):
        return ["本次仅含个人划线与想法。"]
    if any("已点赞他人评论接口读取失败" in w for w in warnings):
        return ["读者观点未能读取，本次仅含个人笔记。"]
    if any("仅核验本次提供" in w for w in warnings):
        return ["读者观点仅含本次指定的想法。"]
    if any("读者观点" in w or "想法" in w or "评论" in w or "扫描" in w for w in warnings):
        return ["读者观点可能不完整。"]
    return list(dict.fromkeys(display_notice(w) for w in warnings))[:1]


def display_units(book: Notebook):
    """每个原文与其相关评论是一个阅读单元，Word 和 Notion 使用同一关联。"""
    yield [("subtitle", f"{book.author} · 读书笔记")]
    for notice in concise_notices(book.warnings):
        yield [("notice", notice)]
    if not book.chapters:
        yield [("paragraph", "暂无可导出的笔记。")]
    for chapter in book.chapters:
        yield [("heading_2", chapter.title)]
        for group in chapter.groups:
            base = group.highlight.text if group.highlight else group.quote
            units: dict[str, list[tuple[str, str]]] = {}
            if base:
                quote = dated_text(group.highlight, book.show_dates) if group.highlight else base
                units[base.strip()] = [("quote", quote)]
            elif not group.thoughts:
                continue
            for note in group.thoughts:
                source = note.quote.strip() if note.quote else base.strip()
                if source not in units:
                    units[source] = [("quote", note.quote)] if source else []
                kind, label = ("thought", "我的想法") if note.kind == "thought" else ("liked", "读者观点")
                units[source].append((kind, label + "\n" + dated_text(note, book.show_dates)))
            yield from units.values()


def entries(book: Notebook):
    """保留平铺语义流供检查使用；布局使用 display_units 保留原文关联。"""
    for unit in display_units(book):
        yield from unit


def safe_filename(title: str) -> str:
    return (re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).strip(" .") or "未命名书籍")[:80]


def export_word(book: Notebook, output: Path) -> Path:
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        raise ExportError("缺少 python-docx，请先运行 python -m pip install -r requirements.txt。") from None
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    for name, size in (("Normal", 11), ("Heading 1", 24), ("Heading 2", 15),
                       ("Quote", 11), ("Subtitle", 10)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.italic = False  # Word 内置 Quote / Subtitle 会继承斜体，需显式关闭。
        italic_cs = OxmlElement("w:iCs")
        italic_cs.set(qn("w:val"), "0")
        style.element.get_or_add_rPr().append(italic_cs)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(10)
        style.paragraph_format.line_spacing = 1.5
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.line_spacing = 1.2
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(24)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(20)
    doc.styles["Subtitle"].paragraph_format.keep_with_next = True
    doc.styles["Quote"].paragraph_format.left_indent = Cm(0.55)
    doc.styles["Quote"].paragraph_format.right_indent = Cm(0.2)
    doc.add_heading(book.title, level=1)
    planned = ((kind, text, index, unit) for unit in display_units(book)
               for index, (kind, text) in enumerate(unit))
    for kind, text, index, unit in planned:
        if kind == "heading_2":
            doc.add_heading(text, level=2)
            continue
        p = doc.add_paragraph(text, style={"quote": "Quote", "subtitle": "Subtitle"}.get(kind, "Normal"))
        p.paragraph_format.widow_control = True
        if kind in ("thought", "liked"):
            p.paragraph_format.left_indent = Cm(0.55)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), {"thought": "F1F6F3", "liked": "F7F4ED"}[kind])
            p._p.get_or_add_pPr().append(shade)
            # 标签单独加粗，正文保持正常字重。
            p.clear()
            label, _, body = text.partition("\n")
            run = p.add_run(label)
            run.bold = True
            run.font.size = Pt(10)
            if body:
                p.add_run("\n" + body)
        elif kind == "notice":
            p.paragraph_format.space_after = Pt(8)
            for run in p.runs:
                run.font.size, run.font.color.rgb = Pt(9), RGBColor.from_string("667085")
        if book.show_dates and kind in ("quote", "thought", "liked") and p.runs:
            last = p.runs[-1]
            tail = re.search(r"（\d{4}-\d{2}-\d{2}）$", last.text)
            if tail:
                last.text = last.text[:tail.start()]
                stamp = p.add_run(tail[0])
                stamp.font.size = Pt(9)
                stamp.font.color.rgb = RGBColor.from_string("667085")
        for run in p.runs:
            run.italic = False
        if kind in ("quote", "thought", "liked"):
            p.paragraph_format.space_after = Pt(16.5)  # 每条笔记后留约一行，正文仍通过缩进/评论关联。
            short_unit = sum(len(value) for _, value in unit) <= 700 and len(unit) <= 6
            p.paragraph_format.keep_with_next = index < len(unit) - 1 and (short_unit or kind == "quote")
            if any(t == "quote" for t, _ in unit):
                # 一组原文与批注共享细左线和缩进，视觉上形成同一阅读单元。
                p.paragraph_format.left_indent = Cm(0.55)
                borders = OxmlElement("w:pBdr")
                left = OxmlElement("w:left")
                for key, value in (("val", "single"), ("sz", "8"), ("space", "10"), ("color", "D8DFDB")):
                    left.set(qn("w:" + key), value)
                borders.append(left)
                p._p.get_or_add_pPr().append(borders)
    doc.core_properties.title = book.title + " 读书笔记"
    doc.core_properties.author = ""
    output.mkdir(parents=True, exist_ok=True)
    path = output / f"《{safe_filename(book.title)}》- 读书笔记.docx"
    # 同名文件不覆盖，采用独占创建避免竞争条件。
    index = 1
    while True:
        try:
            with path.open("xb") as stream:
                doc.save(stream)
            return path.resolve()
        except FileExistsError:
            index += 1
            path = output / f"《{safe_filename(book.title)}》- 读书笔记 ({index}).docx"


def notion_plan(book: Notebook) -> tuple[list[dict], list[dict]]:
    """普通正文 + 章节横线 + 挂在原文段落上的原生评论。

    公共 API 只能按 block_id 关联整段，不能创建新的文字选区讨论。
    聊天连接的 selected-text 评论由连接器单独完成，不能混同这两种能力。
    """
    def make_blocks(kind: str, text: str) -> list[dict]:
        block_type = {"quote": "paragraph", "notice": "paragraph", "subtitle": "paragraph"}.get(kind, kind)
        result = []
        for part in text_chunks(text):
            body: dict = {"rich_text": [{"type": "text", "text": {"content": part},
                                         "annotations": {"italic": False}}]}
            if kind == "notice":
                body["color"] = "gray"
            result.append({"object": "block", "type": block_type, block_type: body})
        return result
    blocks, comments = [], []
    chapter_index, seen_chapter = None, False
    for unit in display_units(book):
        if not unit:
            continue
        anchor_index = None
        for kind, text in unit:
            if kind == "heading_2":
                if seen_chapter:
                    blocks.append({"object": "block", "type": "divider", "divider": {}})
                seen_chapter = True
                chapter_index = len(blocks)
            if kind in ("thought", "liked"):
                anchor = anchor_index if anchor_index is not None else chapter_index
                if anchor is None:
                    raise ExportError("想法缺少可以关联的原文或章节，不能创建独立悬空评论。")
                label, _, body = text.partition("\n")
                parts = text_chunks(body)
                for offset in range(0, len(parts), 20):
                    rich_text = [{"type": "text", "text": {"content": label + ("（续）" if offset else "")},
                                  "annotations": {"bold": True, "italic": False}}]
                    rich_text.extend({"type": "text", "text": {"content": ("\n" if index == 0 else "") + part},
                                      "annotations": {"italic": False}}
                                     for index, part in enumerate(parts[offset:offset + 20]))
                    comments.append({"anchor_index": anchor, "rich_text": rich_text})
                continue
            if kind == "quote":
                anchor_index = len(blocks)
            blocks.extend(make_blocks(kind, text))
            if kind in ("subtitle", "heading_2", "quote"):
                # 真正的空段落；仅 Markdown 换行不会在 Notion 生成留白。
                blocks.append({"object": "block", "type": "paragraph", "paragraph": {"rich_text": []}})
    return blocks, comments


def notion_blocks(book: Notebook) -> list[dict]:
    return notion_plan(book)[0]


def text_chunks(text: str, limit: int = 1800) -> list[str]:
    chunks, current, size = [], [], 0
    for char in text or " ":
        width = 2 if ord(char) > 0xFFFF else 1
        if size + width > limit:
            chunks.append("".join(current))
            current, size = [], 0
        current.append(char)
        size += width
    if current:
        chunks.append("".join(current))
    return chunks


def notion_id(value: str) -> str:
    # 只从路径解析 ID，避免把 URL 中 ?v=视图 ID 当作数据库 ID。
    path = urlsplit(value).path if "://" in value else value.split("?")[0]
    matches = re.findall(r"(?i)[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}|[0-9a-f]{32}", path)
    if not matches:
        raise ExportError("Notion ID 格式不正确，请输入完整页面/数据库 URL 或 32 位 UUID。")
    return str(uuid.UUID(matches[-1]))


def atomic_json(path: Path, value: dict):
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def block_signature(block: dict) -> tuple:
    kind = block["type"]
    return kind, "".join(r.get("plain_text", r.get("text", {}).get("content", ""))
                         for r in block[kind].get("rich_text", []))


def block_count(block: dict) -> int:
    return 1 + sum(block_count(child) for child in block.get(block.get("type"), {}).get("children", []))


def export_notion(book: Notebook, token: str, target: str, target_kind: str, output: Path,
                  client: Any = None) -> str:
    """新建书籍页面，或向指定已有页面追加；检查点同时覆盖正文和原生评论。

    写入网络超时/5xx 不盲目重试，避免重复页/重复块。429/529 明确拒绝才退避。
    若没有本地检查点，重复运行视作新的导出；不会覆盖用户已有页面。
    """
    if client is None:
        try:
            from notion_client import Client
        except ImportError:
            raise ExportError("缺少 notion-client，请先安装 requirements.txt。") from None
        client = Client(auth=token, notion_version="2026-03-11", timeout_ms=30000)
    identifier = notion_id(target)
    append_mode = target_kind == "existing_page"

    def call(func: Callable, **kwargs):
        for attempt in range(4):
            try:
                time.sleep(0.4)
                return func(**kwargs)
            except Exception as exc:
                status = getattr(exc, "status", None)
                if status in (429, 529) and attempt < 3:
                    headers = getattr(exc, "headers", {}) or {}
                    delay = max(2 ** attempt, number(headers.get("retry-after"), 0))
                    print(f"  Notion 限流，{delay} 秒后重试…")
                    time.sleep(delay)
                    continue
                raise ExportError(f"Notion 请求失败（HTTP {status or '未知/网络超时'}）。"
                                  "请检查 Integration 授权和网络；若已创建页面，可重跑并恢复检查点。") from None

    if target_kind == "database":
        database = call(client.databases.retrieve, database_id=identifier)
        sources = database.get("data_sources", [])
        if not sources:
            raise ExportError("数据库没有可访问的数据源，请检查 Integration 的访问权限。")
        if len(sources) > 1:
            print("目标数据库包含多个数据源：")
            for index, source in enumerate(sources, 1):
                print(f"  {index}. {source.get('name', '未命名')} ({source['id']})")
            source = sources[ask_index("请选择要写入的数据源序号：", len(sources))]
        else:
            source = sources[0]
        identifier, target_kind = source["id"], "data_source"
    if target_kind == "data_source":
        source = call(client.data_sources.retrieve, data_source_id=identifier)
        title_properties = [k for k, v in source["properties"].items() if v.get("type") == "title"]
        if len(title_properties) != 1:
            raise ExportError("目标数据源没有唯一的 title 字段。")
        parent = {"type": "data_source_id", "data_source_id": identifier}
        title_property = title_properties[0]
    else:
        target_page = call(client.pages.retrieve, page_id=identifier)
        parent, title_property = {"type": "page_id", "page_id": identifier}, "title"
    blocks, comment_jobs = notion_plan(book)
    if append_mode:
        # 指定已有页面时只追加一个新的书籍区段，不替换已有正文或属性。
        prefix = [{"object": "block", "type": "divider", "divider": {}},
                  {"object": "block", "type": "heading_1", "heading_1": {
                      "rich_text": [{"type": "text", "text": {"content": text_chunks(book.title)[0]}}]}}]
        blocks = prefix + blocks
        comment_jobs = [{**job, "anchor_index": job["anchor_index"] + len(prefix)} for job in comment_jobs]
    fingerprint = hashlib.sha256(json.dumps(["native-comments-v1", target_kind, parent, book.title, blocks, comment_jobs], sort_keys=True,
                                            ensure_ascii=False).encode()).hexdigest()
    output.mkdir(parents=True, exist_ok=True)
    journal = output / f"notion-{fingerprint[:24]}.json"
    state = json.loads(journal.read_text(encoding="utf-8")) if journal.exists() else {}
    if not state.get("page_id"):
        page = ({"id": identifier, "url": target_page.get("url", "")} if append_mode else
                call(client.pages.create, parent=parent,
                     properties={title_property: {"title": [{"text": {"content": text_chunks(book.title + " 读书笔记")[0]}}]}}))
        state = {"page_id": page["id"], "url": page.get("url", ""), "fingerprint": fingerprint}
        atomic_json(journal, state)  # 先落盘 ID，再追加内容
    page_id = state["page_id"]
    print(f"  Notion 页面 ID：{page_id}")
    print(f"  断点文件：{journal.resolve()}")
    def read_children(parent_id):
        result, cursor = [], None
        while True:
            data = call(client.blocks.children.list, block_id=parent_id, page_size=100,
                        **({"start_cursor": cursor} if cursor else {}))
            result.extend(data["results"])
            if not data.get("has_more"):
                return result
            next_cursor = data.get("next_cursor")
            if not next_cursor or next_cursor == cursor:
                raise ExportError("Notion 远端块列表游标停滞，停止恢复。")
            cursor = next_cursor

    def same_block(actual, expected):
        if block_signature(actual) != block_signature(expected):
            return False
        actual_body = actual.get(actual["type"], {})
        expected_children = expected.get(expected["type"], {}).get("children", [])
        if "children" in actual_body:
            actual_children = actual_body["children"]
        elif actual.get("has_children") or expected_children:
            if not actual.get("id"):
                return False
            actual_children = read_children(actual["id"])
        else:
            actual_children = []
        return len(actual_children) == len(expected_children) and all(
            same_block(a, b) for a, b in zip(actual_children, expected_children))

    existing_all = read_children(page_id)
    if append_mode and "base_block_ids" not in state:
        state["base_block_ids"] = [block["id"] for block in existing_all]
        atomic_json(journal, state)
    base_ids = state.get("base_block_ids", [])
    if [block.get("id") for block in existing_all[:len(base_ids)]] != base_ids:
        raise ExportError("指定页面的原有区段已移动或删除，停止恢复以保护页面内容。")
    existing = existing_all[len(base_ids):]
    if len(existing) > len(blocks) or any(not same_block(a, b) for a, b in zip(existing, blocks)):
        raise ExportError("Notion 页面内容已被修改，与本地导出不一致，停止续写以保护现有内容。")
    offset = len(existing)
    while offset < len(blocks):
        # 同时限制数量和 JSON 字节；低于 Notion 的 100 块 / 500 KB 请求上限。
        batch, size, count = [], 0, 0
        for block in blocks[offset:offset + 100]:
            block_size = len(json.dumps(block, ensure_ascii=True).encode())
            nested_count = block_count(block)
            if batch and (size + block_size > 350_000 or count + nested_count > 900):
                break
            batch.append(block)
            size += block_size
            count += nested_count
        call(client.blocks.children.append, block_id=page_id, children=batch)
        offset += len(batch)
        state["written_blocks"] = offset
        atomic_json(journal, state)
        print(f"  已写入 {offset}/{len(blocks)} 个块")
    # 重新读取实际 block ID 后，把想法和观点作为原生评论挂到对应原文。
    if comment_jobs:
        saved_blocks = read_children(page_id)[len(base_ids):]
        if len(saved_blocks) != len(blocks) or any(not same_block(a, b) for a, b in zip(saved_blocks, blocks)):
            raise ExportError("正文写入后与预期不一致，停止创建评论，避免关联到错误原文。")
        receipts = state.setdefault("comment_receipts", {})
        remote_comments = {}
        for index, job in enumerate(comment_jobs):
            block_id = saved_blocks[job["anchor_index"]].get("id")
            if not block_id:
                raise ExportError("正文已写入，但未取得原文 block ID，无法关联评论。")
            target_text = "".join(part["text"]["content"] for part in job["rich_text"])
            receipt_key = hashlib.sha256(json.dumps([index, block_id, job["rich_text"]], sort_keys=True).encode()).hexdigest()
            if receipt_key in receipts:
                continue  # 用户已编辑或解决的评论不自动重建。
            try:
                if block_id not in remote_comments:
                    found, cursor = [], None
                    while True:
                        response = call(client.comments.list, block_id=block_id, page_size=100,
                                        **({"start_cursor": cursor} if cursor else {}))
                        found.extend(response["results"])
                        if not response.get("has_more"):
                            break
                        next_cursor = response.get("next_cursor")
                        if not next_cursor or next_cursor == cursor:
                            raise ExportError("评论分页游标没有前进。")
                        cursor = next_cursor
                    remote_comments[block_id] = found
                match = next((comment for comment in remote_comments[block_id] if comment.get("id") not in receipts.values() and "".join(
                    part.get("plain_text", part.get("text", {}).get("content", ""))
                    for part in comment.get("rich_text", [])) == target_text), None)
                if match is None:
                    match = call(client.comments.create, parent={"block_id": block_id}, rich_text=job["rich_text"])
                    remote_comments[block_id].append(match)
                receipts[receipt_key] = match["id"]
                atomic_json(journal, state)
                print(f"  已关联原生评论 {index + 1}/{len(comment_jobs)} 条")
            except ExportError as exc:
                raise ExportError("Notion 正文已写入，但原生评论未完成。请检查 Integration 的读取评论和插入评论权限；"
                                  f"修复后用相同内容重跑，可恢复进度。详情：{exc}") from None
    return state.get("url") or f"https://www.notion.so/{page_id.replace('-', '')}"


def ask(prompt: str) -> str:
    value = input(prompt).strip()
    if value.lower() in ("q", "quit", "exit"):
        raise KeyboardInterrupt
    return value


def read_clipboard() -> str:
    """用户选择后才读取 Windows Unicode 文本，不启动 shell、不打印或写入凭证。

    ctypes 使用明确的指针类型，避免 64 位 Windows 的句柄被截断。
    读取后保留剪贴板原样，便于用户重试；并不扫描浏览器凭证存储。
    """
    if os.name != "nt":
        raise ExportError("从剪贴板读取目前仅支持 Windows，请选择隐藏输入。")
    import ctypes
    from ctypes import wintypes

    user32 = ctypes.WinDLL("user32", use_last_error=True)
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    user32.OpenClipboard.argtypes = [wintypes.HWND]
    user32.OpenClipboard.restype = wintypes.BOOL
    user32.CloseClipboard.argtypes = []
    user32.CloseClipboard.restype = wintypes.BOOL
    user32.GetClipboardData.argtypes = [wintypes.UINT]
    user32.GetClipboardData.restype = wintypes.HANDLE
    kernel32.GlobalLock.argtypes = [wintypes.HANDLE]
    kernel32.GlobalLock.restype = ctypes.c_void_p
    kernel32.GlobalUnlock.argtypes = [wintypes.HANDLE]
    kernel32.GlobalUnlock.restype = wintypes.BOOL
    kernel32.GlobalSize.argtypes = [wintypes.HANDLE]
    kernel32.GlobalSize.restype = ctypes.c_size_t
    opened = False
    for _ in range(5):
        if user32.OpenClipboard(None):
            opened = True
            break
        time.sleep(0.1)
    if not opened:
        raise ExportError("剪贴板暂时被占用，请关闭复制菜单后重试。")
    try:
        handle = user32.GetClipboardData(13)  # CF_UNICODETEXT
        if not handle:
            raise ExportError("剪贴板中没有文本，请先复制 Cookie 的完整值。")
        size = kernel32.GlobalSize(handle)
        if size < 2 or size > 2_000_000 or size % 2:
            raise ExportError("剪贴板文本大小异常，请只复制所需凭证。")
        pointer = kernel32.GlobalLock(handle)
        if not pointer:
            raise ExportError("无法读取剪贴板，请重新复制后重试。")
        try:
            # 按分配长度读取，避免扫描越界；Windows Unicode 文本以 NUL 结尾。
            return ctypes.string_at(pointer, size).decode("utf-16-le").split("\0", 1)[0]
        finally:
            kernel32.GlobalUnlock(handle)
    except (UnicodeError, OSError):
        raise ExportError("读取剪贴板文本失败，请重新复制后重试。") from None
    finally:
        user32.CloseClipboard()


def secret(prompt: str, env: str, validator: Callable[[str], Any] | None = None) -> str:
    value = os.environ.get(env, "").strip()
    if value:
        if validator:
            try:
                validator(value)
            except ExportError as exc:
                raise ExportError(f"环境变量 {env} 无效：{exc} 请更新或清除该变量后重试。") from None
        print(f"  已读取环境变量 {env}")
        return value
    print(prompt.replace("（隐藏输入）", "").rstrip("：:"))
    while True:
        try:
            if os.name == "nt":
                method = choice("输入方式：1. 从剪贴板读取（推荐）；2. 隐藏输入 [默认 1]：", ("1", "2"), "1")
            else:
                method = "2"
            if method == "1":
                ask("请先在浏览器中复制凭证完整值，回到这里按回车读取（无需粘贴）：")
                value = read_clipboard().strip()
            else:
                value = getpass.getpass("请输入凭证（隐藏输入）：").strip()
            if not value:
                raise ExportError("输入不能为空，请重新复制或输入。")
            if validator:
                validator(value)
            status = "，Cookie 必需字段检查通过（登录有效性待接口验证）。" if validator is parse_cookie else (
                "，格式检查通过（授权有效性待接口验证）。" if validator else "。")
            print("  已读取凭证" + status)
            return value
        except ExportError as exc:
            print(f"  {exc}")
            print("  请重试，或输入 q / 按 Ctrl+C 退出。")


def choice(prompt: str, options: tuple[str, ...], default: str = "") -> str:
    while True:
        value = ask(prompt) or default
        if value in options:
            return value
        print("  请输入 " + " / ".join(options) + "。")


def ask_index(prompt: str, count: int) -> int:
    while True:
        value = ask(prompt)
        if value.isdigit() and 1 <= int(value) <= count:
            return int(value) - 1
        print(f"  请输入 1 到 {count} 之间的序号。")


def ask_notion_destination() -> tuple[str, str]:
    """导入前明确选择新建笔记或指定位置，指定已有页面仅追加。"""
    mode = choice("Notion 导入方式：1. 新建笔记；2. 导入到指定位置：", ("1", "2"))
    if mode == "1":
        print("  新建独立书籍笔记。独立脚本需选择已授权的保存容器；聊天连接可直接新建私有笔记。")
        kind = choice("新笔记保存到：1. 数据库；2. 父页面下；3. Data Source：", ("1", "2", "3"))
        target_kind = {"1": "database", "2": "page", "3": "data_source"}[kind]
    else:
        kind = choice("指定位置：1. 追加到已有笔记页面；2. 指定父页面下新建；3. 指定数据库：", ("1", "2", "3"))
        target_kind = {"1": "existing_page", "2": "page", "3": "database"}[kind]
        if target_kind == "existing_page":
            print("  将在目标页面末尾追加本书笔记，保留原有正文和属性。")
    target = os.environ.get("NOTION_TARGET_ID", "").strip() or ask("请输入目标页面/数据库的链接或 ID：")
    notion_id(target)
    return target_kind, target


def select_book(books: list[dict]) -> dict:
    if not books:
        raise ExportError("当前书架没有可导出的书籍。")
    print(f"\n我的书架（{len(books)} 本可选书籍）：")
    for i, book in enumerate(books, 1):
        print(f"  {i:>3}. {clean(book.get('title') or '未命名')} — {clean(book.get('author'))}"
              f"  [bookId={book['bookId']}]")
    while True:
        value = ask("\n请问您需要对‘我的书架’中的哪本书进行笔记总结？（序号/关键词）：")
        if value.isdigit() and 1 <= int(value) <= len(books):
            return books[int(value) - 1]
        if not value:
            continue
        matches = [b for b in books if value.casefold() in str(b.get("title", "")).casefold()]
        if len(matches) == 1:
            return matches[0]
        if not matches:
            print("  没有匹配的书籍，请重新输入。")
            continue
        for index, book in enumerate(matches, 1):
            print(f"  {index}. {book.get('title')} — {book.get('author', '')} [bookId={book['bookId']}]")
        return matches[ask_index("匹配到多本书，请输入上述匹配列表中的序号：", len(matches))]


def demo_book(timezone: ZoneInfo) -> Notebook:
    # 演示数据完全离线，不会访问微信账号或 Notion。
    return build_notebook({"bookId": "demo", "title": "阅读与思考", "author": "示例作者"},
        [{"chapterUid": 19, "chapterIdx": 1, "title": "第一章 阅读的起点"},
         {"chapterUid": 3, "chapterIdx": 2, "title": "第二章 把思考变成行动"}],
        [{"bookmarkId": "h1", "chapterUid": 19, "range": "10-30", "markText": "阅读的价值，在于让我们重新看见熟悉的世界。", "createTime": 1793000000},
         {"bookmarkId": "h2", "chapterUid": 3, "markText": "把一条启发写下来，再为它安排一个可以执行的小行动。", "createTime": 1793000500},
         {"bookmarkId": "h3", "chapterUid": 999, "markText": "没有目录位置的摘录，也值得被妥善保存。"}],
        [{"reviewId": "m1", "chapterUid": 19, "range": "10-30", "content": "读完这一段，我想为每天的阅读留出一页手写记录。", "createTime": 1793000200}],
        [{"reviewId": "r1", "chapterUid": 19, "range": "10-30", "content": "好的阅读不是收集结论，而是学习提出自己的问题。", "author": {"name": "一位读者"}, "likesCount": 128, "createTime": 1793000300}],
        timezone, ["这是离线示例数据，用于预览导出版式。"])


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--demo", action="store_true", help="使用演示数据导出 Word，不联网")
    parser.add_argument("--check-review", nargs="?", const="", default=None, metavar="REVIEW_ID_OR_URL",
                        help="只验证单条想法是否返回点赞状态，不抓书架、不导出正文")
    parser.add_argument("--personal-only", action="store_true", help="仅导出个人划线与想法，使用 Agent Key 时无需 Cookie，不抓取点赞评论")
    parser.add_argument("--liked-source", choices=("paragraphs", "links"), default="paragraphs",
                        help="读者观点来源：paragraphs 自动扫描可发现段落；links 仅核验指定链接")
    parser.add_argument("--review-id", action="append", default=[], metavar="ID_OR_URL",
                        help="优先核验的想法 ID 或链接，可重复；仅收录仍由当前账号点赞且属于所选书籍的想法")
    parser.add_argument("--max-review-checks", type=int, default=200, help="官方单条想法核验上限，默认 200")
    parser.add_argument("--max-review-ranges", type=int, default=500, help="自动扫描原文范围上限，默认 500")
    dates = parser.add_mutually_exclusive_group()
    dates.add_argument("--show-dates", dest="show_dates", action="store_true", help="在正文后标注简短日期，如（2026-10-28）")
    dates.add_argument("--hide-dates", dest="show_dates", action="store_false", help="不在导出内容中显示日期")
    parser.set_defaults(show_dates=None)
    parser.add_argument("--output", type=Path, default=Path("Weread_Notes_Output"), help="输出文件夹")
    parser.add_argument("--timezone", default="Asia/Shanghai", help="日期时区，默认 Asia/Shanghai")
    parser.add_argument("--min-likes", type=int, default=0, help="已点赞评论的额外点赞数门槛，默认不设门槛")
    parser.add_argument("--max-pages", type=int, default=1000, help="翻页保护上限，达到时明确报错")
    args = parser.parse_args()
    if args.min_likes < 0 or min(args.max_pages, args.max_review_checks, args.max_review_ranges) < 1:
        parser.error("--min-likes 必须 >= 0，所有扫描上限必须 >= 1")
    if args.personal_only and args.review_id:
        parser.error("--personal-only 不抓取读者观点，请移除此选项后使用 --review-id。")
    if args.liked_source == "links" and not args.review_id:
        parser.error("--liked-source links 需要至少一个 --review-id。")
    try:
        if args.check_review is not None:
            reference = args.check_review or ask("请输入一条已点赞想法的详情链接或 reviewId：")
            review_id_from_reference(reference)  # 先检查引用格式，再读取凭证
            key = secret("请输入 Agent API Key（隐藏输入）：", "WEREAD_API_KEY", validator=validate_api_key)
            probe_review(Weread(api_key=key), reference)
            return 0
        timezone = ZoneInfo(args.timezone)
        for reference in args.review_id:
            review_id_from_reference(reference)
        if args.demo:
            print("正在排版离线演示文档…")
            preview = demo_book(timezone)
            preview.show_dates = bool(args.show_dates)
            path = export_word(preview, args.output)
            print(f"导出成功：{path}")
            return 0
        print("微信读书笔记导出器（输入 q 或按 Ctrl+C 退出）")
        print("[1/5] 凭证初始化：凭证仅保留在本次进程内存中。")
        if not args.personal_only:
            print("  读者观点模式：" + ("只核验指定链接。" if args.liked_source == "links" else "扫描可发现段落，指定链接优先核验。"))
        mode = choice("认证方式：1. Cookie；2. Agent API Key [默认 1]：", ("1", "2"), "1")
        api_key = secret("请输入 Agent API Key（隐藏输入）：", "WEREAD_API_KEY", validator=validate_api_key) if mode == "2" else ""
        if api_key and not args.personal_only:
            print("  本次使用官方段落想法与单条详情接口核验点赞，不需要 Cookie。")
        if args.personal_only:
            print("  本次仅导出个人划线与想法，不抓取已点赞他人评论。")
        cookie = "" if api_key else secret(
            "请输入微信读书完整 Cookie（隐藏输入）：", "WEREAD_COOKIE", validator=parse_cookie)
        client = Weread(cookie, api_key, args.max_pages, interactive_auth=True, personal_only=args.personal_only,
                        liked_source=args.liked_source, review_ids=args.review_id,
                        max_review_checks=args.max_review_checks, max_review_ranges=args.max_review_ranges)
        print("[2/5] 正在获取我的书架…")
        book = select_book(client.shelf())
        print(f"\n已选择：《{book.get('title')}》 [bookId={book['bookId']}]")
        target = choice("请选择导出方式：1. 导出为本地精美 Word (.docx) 文档；2. 同步到 Notion 数据库：", ("1", "2"))
        token = target_id = target_kind = ""
        if target == "2":
            target_kind, target_id = ask_notion_destination()
            token = secret("请输入您的 Notion Integration Token（隐藏输入）：", "NOTION_TOKEN")
        if args.show_dates is None:
            show_dates = choice("笔记和评论是否显示记录日期？1. 显示（正文后标注日期）；2. 不显示 [默认 2]：", ("1", "2"), "2") == "1"
        else:
            show_dates = args.show_dates
        client.show_dates = show_dates
        print("  导出格式：" + ("正文后标注日期，如（2026-10-28）。" if show_dates else "不显示记录日期。"))
        notebook = client.collect(book, timezone, args.min_likes)
        for warning in notebook.warnings:
            print("  提示：" + warning)
        highlights = sum(bool(g.highlight) for c in notebook.chapters for g in c.groups)
        thoughts = sum(n.kind == "thought" for c in notebook.chapters for g in c.groups for n in g.thoughts)
        liked = sum(n.kind == "liked" for c in notebook.chapters for g in c.groups for n in g.thoughts)
        print(f"[4/5] 已分类：{highlights} 条划线、{thoughts} 条个人想法、{liked} 条已点赞他人评论。")
        print("[4/5] 正在排版…" if target == "1" else "[4/5] 正在写入 Notion…")
        result = export_word(notebook, args.output) if target == "1" else export_notion(
            notebook, token, target_id, target_kind, args.output)
        print(f"[5/5] 导出成功：{result}")
        return 0
    except ZoneInfoNotFoundError:
        print("错误：时区无效或缺少时区数据库，请安装 tzdata 并检查 --timezone。", file=sys.stderr)
    except (KeyboardInterrupt, EOFError):
        print("\n已取消。")
        return 130
    except ExportError as exc:
        print(f"错误：{exc}", file=sys.stderr)
    except OSError:
        print("错误：无法读写输出目录，请检查文件占用、磁盘空间及权限。", file=sys.stderr)
    except Exception as exc:
        # 不输出原始异常/请求体，避免第三方库异常携带 Token 或 Cookie。
        print(f"错误：出现未预期的 {type(exc).__name__}，未能完成导出。", file=sys.stderr)
    return 1


if __name__ == "__main__":
    sys.exit(main())
