import copy
import json
from pathlib import Path
import sys
import subprocess
import tempfile
import unittest

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / 'plugins/weread-notes/scripts'))
import bridge
import incremental as inc
import weread_notes as core


def book(extra=False, show_dates=False):
    first = core.Note('highlight', '1', '旧原文', '2026-01-01', '1-5', 'mark1')
    old = core.Note('thought', '1', '旧想法', '2026-01-01', '1-5', 'review1', '旧原文')
    thoughts = [old]
    if extra:
        thoughts.append(core.Note('liked', '1', '新的读者观点', '2026-02-01', '1-5', 'review2', '旧原文'))
    groups = [core.Group(first, thoughts)]
    if extra:
        groups.append(core.Group(core.Note('highlight', '1', '新增原文', '2026-02-01', '20-30', 'mark2')))
    return core.Notebook('book-1', '增量测试', '作者', [core.Chapter('1', '第一章', 0, groups)], [], show_dates)


def snapshot():
    return {'book_id': 'book-1', 'page_id': 'page-1', 'body_complete': True, 'comments_complete': True,
            'chapters': [{'title': '第一章', 'paragraphs': [{'text': '旧原文', 'id': 'block1'}]}],
            'comments': [{'id': 'comment1', 'chapter_title': '第一章', 'selection': '旧原文', 'text': '我的想法\n旧想法'}]}


class IncrementalTests(unittest.TestCase):
    def test_record_dates_are_plain_numeric_text_in_body_and_comments(self):
        current = book(show_dates=True)
        plan = bridge.connector_plan(current)
        self.assertIn('（2026-01-01）', plan['content'])
        self.assertNotIn('<date ', plan['content'])
        self.assertNotIn('@May', plan['content'])
        for comment in plan['comments']:
            self.assertTrue(all(part['type'] == 'text' for part in comment['rich_text']))
            self.assertIn('（2026-01-01）', ''.join(p['text']['content'] for p in comment['rich_text']))
        blocks, comments = core.notion_plan(current)
        for block in blocks:
            body = block.get(block['type'], {})
            self.assertTrue(all(p['type'] == 'text' for p in body.get('rich_text', [])))
        self.assertTrue(all(p['type'] == 'text' for c in comments for p in c['rich_text']))

    def test_only_new_quote_and_comment_are_added(self):
        original = snapshot()
        before = copy.deepcopy(original)
        delta = inc.notion_delta(bridge.connector_plan(book(extra=True)), original)
        self.assertEqual((delta['added_sources'], delta['added_comments'], delta['retained']), (1, 1, 2))
        self.assertEqual(delta['comments'][0]['selection'], '旧原文')
        self.assertEqual(delta['comments'][0]['anchor_block_id'], 'block1')
        self.assertEqual(original, before)

    def test_resync_after_success_is_noop(self):
        plan = bridge.connector_plan(book(extra=True))
        remote = snapshot()
        remote['chapters'][0]['paragraphs'].append({'text': '新增原文', 'id': 'block2'})
        remote['comments'].append({'chapter_title': '第一章', 'selection': '旧原文', 'text': '读者观点\n新的读者观点'})
        delta = inc.notion_delta(plan, remote)
        self.assertEqual((delta['added_sources'], delta['added_comments']), (0, 0))
        self.assertEqual(delta['retained'], 4)

    def test_date_visibility_does_not_create_duplicates(self):
        remote = snapshot()
        remote['chapters'][0]['paragraphs'][0]['text'] += '（2026-01-01）'
        remote['comments'][0]['text'] += '（2026-01-01）'
        delta = inc.notion_delta(bridge.connector_plan(book(show_dates=False)), remote)
        self.assertEqual((delta['added_sources'], delta['added_comments']), (0, 0))

    def test_same_quote_in_different_chapter_is_not_deduplicated(self):
        current = book()
        current.chapters.append(core.Chapter('2', '第二章', 1, [core.Group(core.Note('highlight', '2', '旧原文', '', '2-6', 'mark2'))]))
        delta = inc.notion_delta(bridge.connector_plan(current), snapshot())
        self.assertEqual(delta['added_sources'], 1)
        self.assertEqual(delta['chapters'][0]['title'], '第二章')

    def test_manual_edits_preserved_and_new_comment_uses_current_anchor(self):
        plan = bridge.connector_plan(book(extra=True))
        remote = snapshot()
        remote['chapters'][0]['paragraphs'][0]['text'] = '用户改写过的原文'
        remote['comments'][0]['text'] = '用户改写过的想法'
        receipt = {'book_id': 'book-1', 'page_id': 'page-1', 'completed': [
            {'key': plan['records'][0]['key'], 'remote_id': 'block1'},
            {'key': plan['records'][1]['key'], 'remote_id': 'comment1'}]}
        delta = inc.notion_delta(plan, remote, receipt)
        self.assertEqual(delta['added_sources'], 1)
        self.assertEqual(delta['added_comments'], 1)
        self.assertEqual(delta['comments'][0]['selection'], '用户改写过的原文')

    def test_missing_anchor_does_not_duplicate_original_or_attach_wrongly(self):
        plan = bridge.connector_plan(book(extra=True))
        remote = snapshot()
        remote['chapters'][0]['paragraphs'] = []
        receipt = {'book_id': 'book-1', 'page_id': 'page-1', 'completed': [{'key': plan['records'][0]['key']}]}
        delta = inc.notion_delta(plan, remote, receipt, respect_deletions=True)
        self.assertEqual(delta['added_sources'], 1)  # 用户明确保留删除状态时不重建。
        self.assertTrue(delta['unresolved'])

    def test_partial_long_comment_retries_only_unfinished_parts(self):
        current = book()
        current.chapters[0].groups[0].thoughts[0].text = '长评论' * 20000
        plan = bridge.connector_plan(current)
        key = next(r['key'] for r in plan['records'] if r['kind'] == 'thought')
        receipt = {'book_id': 'book-1', 'page_id': 'page-1', 'completed': [],
                   'comment_parts': [{'key': key, 'part': 0, 'remote_id': 'comment-part1'}]}
        remote = snapshot()
        first = inc.comment_parts(next(r for r in plan['records'] if r['key'] == key))[0]
        remote['comments'].append({'id': 'comment-part1', 'chapter_title': '第一章', 'selection': '旧原文',
                                   'text': ''.join(p['text']['content'] for p in first['rich_text'])})
        delta = inc.notion_delta(plan, remote, receipt)
        self.assertTrue(delta['comments'])
        self.assertTrue(all(job['part'] > 0 for job in delta['comments'][0]['jobs']))

    def test_deleted_original_is_restored_despite_old_receipt(self):
        plan = bridge.connector_plan(book())
        remote = snapshot()
        remote['chapters'][0]['paragraphs'] = []
        remote['comments'] = []
        receipt = {'book_id': 'book-1', 'page_id': 'page-1', 'completed': [
            {'key': plan['records'][0]['key'], 'remote_id': 'deleted-block'},
            {'key': plan['records'][1]['key'], 'remote_id': 'deleted-comment'}]}
        delta = inc.notion_delta(plan, remote, receipt)
        self.assertEqual((delta['added_sources'], delta['added_comments']), (1, 1))
        self.assertEqual((delta['restored_sources'], delta['restored_comments']), (1, 1))
        self.assertEqual(delta['comments'][0]['selection'], '旧原文')

    def test_deleted_comment_is_restored_without_duplicating_original(self):
        plan = bridge.connector_plan(book())
        remote = snapshot()
        remote['comments'] = []
        receipt = {'book_id': 'book-1', 'page_id': 'page-1', 'completed': [
            {'key': plan['records'][1]['key'], 'remote_id': 'deleted-comment'}]}
        delta = inc.notion_delta(plan, remote, receipt)
        self.assertEqual((delta['added_sources'], delta['added_comments']), (0, 1))
        self.assertEqual(delta['comments'][0]['anchor_block_id'], 'block1')

    def test_spacing_is_added_once_and_incomplete_body_is_rejected(self):
        plan = bridge.connector_plan(book())
        remote = snapshot()
        self.assertEqual(len(inc.notion_delta(plan, remote)['spacing']), 1)
        remote['chapters'][0]['paragraphs'][0]['blank_after'] = True
        self.assertEqual(inc.notion_delta(plan, remote)['spacing'], [])
        remote['body_complete'] = False
        with self.assertRaises(inc.SyncError):
            inc.notion_delta(plan, remote)

    def test_deleted_long_comment_part_is_not_trusted_from_receipt(self):
        current = book()
        current.chapters[0].groups[0].thoughts[0].text = '长评论' * 20000
        plan = bridge.connector_plan(current)
        key = next(r['key'] for r in plan['records'] if r['kind'] == 'thought')
        receipt = {'book_id': 'book-1', 'page_id': 'page-1',
                   'comment_parts': [{'key': key, 'part': 0, 'remote_id': 'gone'}]}
        delta = inc.notion_delta(plan, snapshot(), receipt)
        self.assertIn(0, [job['part'] for job in delta['comments'][0]['jobs']])

    def test_completed_long_comment_with_missing_part_is_repaired(self):
        current = book()
        current.chapters[0].groups[0].thoughts[0].text = '长评论' * 20000
        plan = bridge.connector_plan(current)
        record = next(r for r in plan['records'] if r['kind'] == 'thought')
        first = inc.comment_parts(record)[0]
        remote = snapshot()
        remote['comments'] = [{'id': 'part0', 'chapter_title': '第一章', 'selection': '旧原文',
                               'text': ''.join(p['text']['content'] for p in first['rich_text'])}]
        receipt = {'book_id': 'book-1', 'page_id': 'page-1',
                   'completed': [{'key': record['key'], 'remote_id': 'part0'}],
                   'comment_parts': [{'key': record['key'], 'part': 0, 'remote_id': 'part0'},
                                     {'key': record['key'], 'part': 1, 'remote_id': 'deleted'}]}
        delta = inc.notion_delta(plan, remote, receipt)
        self.assertEqual([j['part'] for j in delta['comments'][0]['jobs']], [1])

    def test_incomplete_comments_and_wrong_receipt_stop_sync(self):
        plan = bridge.connector_plan(book())
        remote = snapshot()
        remote['comments_complete'] = False
        with self.assertRaises(inc.SyncError):
            inc.notion_delta(plan, remote)
        with self.assertRaises(inc.SyncError):
            inc.notion_delta(plan, snapshot(), {'book_id': 'other', 'page_id': 'page-1'})

    def test_sync_plan_cli_produces_only_additions_without_auth(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'plan.json').write_text(json.dumps(bridge.connector_plan(book(extra=True))), encoding='utf-8')
            (root / 'remote.json').write_text(json.dumps(snapshot()), encoding='utf-8')
            run = subprocess.run([sys.executable, str(ROOT / 'plugins/weread-notes/scripts/launch.py'),
                                  'sync-plan', '--plan', str(root / 'plan.json'), '--existing', str(root / 'remote.json'),
                                  '--output', str(root / 'delta.json')], capture_output=True, encoding='utf-8')
            self.assertEqual(run.returncode, 0, run.stderr + run.stdout)
            summary = json.loads(run.stdout)
            self.assertEqual((summary['added_sources'], summary['added_comments']), (1, 1))
            self.assertTrue((root / 'delta.json').is_file())

    def test_equal_comment_bodies_at_different_sources_stay_associated(self):
        from docx import Document
        current = book()
        current.chapters[0].groups.append(core.Group(
            core.Note('highlight', '1', '另一个原文', '', '40-50', 'mark3'),
            [core.Note('thought', '1', '旧想法', '', '40-50', 'review3', '另一个原文')]))
        delta = inc.notion_delta(bridge.connector_plan(current), snapshot())
        self.assertEqual(delta['added_comments'], 1)
        self.assertEqual(delta['comments'][0]['selection'], '另一个原文')
        with tempfile.TemporaryDirectory() as folder:
            first = inc.export_word_incremental(book(), Path(folder))
            updated = inc.export_word_incremental(current, Path(folder))
            texts = [p.text for p in Document(updated['path']).paragraphs]
            self.assertEqual(texts.count('我的想法\n旧想法'), 2)
            self.assertEqual(updated['added'], 2)

    def test_word_repeated_export_preserves_same_document(self):
        with tempfile.TemporaryDirectory() as folder:
            first = inc.export_word_incremental(book(), Path(folder))
            before = Path(first['path']).read_bytes()
            again = inc.export_word_incremental(book(), Path(folder))
            self.assertEqual(again['path'], first['path'])
            self.assertEqual(again['added'], 0)
            self.assertEqual(again['retained'], 2)
            self.assertEqual(Path(again['path']).read_bytes(), before)

    def test_word_adds_new_items_and_preserves_manual_text_and_tables(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            first = inc.export_word_incremental(book(), Path(folder))
            doc = Document(first['path'])
            source = next(p for p in doc.paragraphs if p.text == '旧原文')
            source.runs[0].text = '手工改写原文'
            doc.add_paragraph('我自己写的补充内容')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = '手工表格'
            doc.save(first['path'])
            updated = inc.export_word_incremental(book(extra=True), Path(folder))
            doc = Document(updated['path'])
            texts = [p.text for p in doc.paragraphs]
            self.assertEqual(updated['added'], 2)
            self.assertIn('手工改写原文', texts)
            self.assertNotIn('旧原文', texts)
            self.assertIn('我自己写的补充内容', texts)
            self.assertEqual(doc.tables[0].cell(0, 0).text, '手工表格')
            self.assertEqual(texts.count('新增原文'), 1)
            self.assertEqual(sum(t.startswith('读者观点\n新的读者观点') for t in texts), 1)
            self.assertLess(texts.index('手工改写原文'), next(i for i, t in enumerate(texts) if t.startswith('读者观点\n')))
            self.assertTrue(Path(updated['path'] + '.before-sync.bak').exists())
            self.assertEqual(inc.export_word_incremental(book(extra=True), Path(folder))['added'], 0)

    def test_legacy_word_without_receipt_is_matched_without_duplicate_text(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            original = core.export_word(book(), Path(folder))
            updated = inc.export_word_incremental(book(extra=True), Path(folder))
            self.assertEqual(Path(updated['path']), original)
            self.assertEqual([p.text for p in Document(original).paragraphs].count('旧原文'), 1)
            self.assertEqual(updated['added'], 2)

    def test_multiple_old_word_files_require_target_in_chat(self):
        with tempfile.TemporaryDirectory() as folder:
            core.export_word(book(), Path(folder))
            core.export_word(book(), Path(folder))
            with self.assertRaises(inc.SyncError) as raised:
                inc.export_word_incremental(book(extra=True), Path(folder))
            self.assertEqual(raised.exception.code, 'word_target_required')
            self.assertEqual(len(raised.exception.choices), 2)


if __name__ == '__main__':
    unittest.main()
